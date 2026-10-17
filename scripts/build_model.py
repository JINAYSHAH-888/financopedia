#!/usr/bin/env python3
"""
Financopedia model builder.

Generates fully formula-driven workbooks from a JSON spec. Every number
outside INPUTS and HISTORICALS is an Excel formula: change one input cell and
the whole workbook moves.

    python3 build_model.py SPEC.json --type dcf        -o TICK-dcf.xlsx
    python3 build_model.py SPEC.json --type statements -o TICK-statements.xlsx
    python3 build_model.py SPEC.json --type lbo        -o TICK-lbo.xlsx
    python3 build_model.py SPEC.json --type ipo        -o TICK-ipo.xlsx

    python3 build_model.py SPEC.json --scan            # rank value drivers
    python3 build_model.py SPEC.json --breakeven wacc,ebit_margin
    python3 build_model.py SPEC.json --history         # what the history says
"""

from __future__ import annotations

import argparse
import os
import sys

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import dcf_core as core  # noqa: E402
from xl_style import (BOX, CUR, F_ESTIMATE, F_FORMULA, F_INPUT, F_NOTE,  # noqa: E402
                      F_SUB, F_TITLE, FILL_BAND, FILL_EST, FILL_OUT, MULT,
                      NUM, PCT, checks_sheet, col, head, label, note, put,
                      readme_sheet, widths, year_widths)

# --------------------------------------------------------------------------
# cell map -- rows are fixed, columns follow the spec's horizons
# --------------------------------------------------------------------------

CELLS = {
    "INPUTS": {
        "tax_rate": "B4", "risk_free": "B5", "equity_risk_premium": "B6",
        "beta": "B7", "cost_of_debt": "B8", "debt_weight": "B9",
        "wacc": "B10", "terminal_growth": "B11", "exit_multiple": "B12",
        "net_debt": "B13", "shares_out": "B14", "current_price": "B15",
        "wacc_step": "B16", "g_step": "B17",
        "first_period_fraction": "B18", "scenario_growth_delta": "B19",
        "vector_header": 21,
        "rev_growth_row": 22, "ebit_margin_row": 23, "da_pct_rev_row": 24,
        "capex_pct_rev_row": 25, "nwc_pct_rev_row": 26,
        "comps_header_row": 29, "comps_start_row": 30,
    },
    "HISTORICALS": {
        "year_row": 3, "revenue": 4, "ebit": 5, "ebit_margin": 6, "da": 7,
        "capex": 8, "nwc_balance": 9, "delta_nwc": 10, "fcf": 11,
        "ppe_net": 12,
        "stats_header": 15, "cagr": 16, "margin_avg": 17, "margin_min": 18,
        "margin_max": 19, "capex_avg": 20, "down_years": 21,
    },
    "PROJECTION": {
        "year_row": 3, "revenue": 4, "ebit": 5, "tax": 6, "nopat": 7, "da": 8,
        "capex": 9, "nwc_balance": 10, "delta_nwc": 11, "fcf": 12,
    },
    "VALUATION": {
        "fcf": 4, "period": 5, "discount_factor": 6, "pv_fcf": 7,
        "pv_explicit": "B9", "terminal_value": "B10", "pv_terminal": "B11",
        "enterprise_value": "B12", "net_debt": "B13", "equity_value": "B14",
        "shares": "B15", "value_per_share": "B16", "current_price": "B17",
        "implied_vs_price": "B18",
        "terminal_ebitda": "B21", "tv_exit": "B22", "pv_tv_exit": "B23",
        "ev_exit": "B24", "value_per_share_exit": "B25",
        "peer_median_ev_ebitda": "B28", "ntm_ebitda": "B29",
        "implied_ev_comps": "B30", "value_per_share_comps": "B31",
    },
    "SENSITIVITY": {
        "x_row": 4, "y_start": 5, "grid_first_col": 3, "helper_start": 14,
    },
}

GRID = 5



def _wacc_refs(spec):
    """Row numbers on the WACC tab that INPUTS links to, so one number has one home."""
    import build_dcf_extras as ex
    w = spec["wacc_build"]
    R = ex.erp_rows(ex.wacc_rows(len(w.get("peers", []))),
                    len(w.get("market_returns", {}).get("years", [])))
    return {"risk_free": R["risk_free"], "equity_risk_premium": R["erp"],
            "beta": R["beta_used"], "cost_of_debt": R["pre_tax_kd"],
            "debt_weight": R["debt_weight"], "wacc": R["wacc"]}


def basis_col(n_vec: int) -> int:
    """Column holding basis / source text, just right of the widest block."""
    return max(2 + n_vec, 5)


# --------------------------------------------------------------------------
# DCF workbook
# --------------------------------------------------------------------------


def build_inputs(wb, spec):
    ws = wb.create_sheet("INPUTS")
    i = spec["inputs"]
    basis, est = i.get("basis", {}), set(i.get("estimates", []))
    nf = core.n_fcst(spec)
    bcol = basis_col(nf)
    I = CELLS["INPUTS"]

    ws["A1"] = "INPUTS — the only assumptions in the model"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue = input. Amber italic = estimated. Black = calculated "
                "from the cells above.")
    ws["A2"].font = F_NOTE

    head(ws, 3, ["Assumption", "Value", "Unit"])
    ws.cell(row=3, column=bcol, value="Basis / source").font = F_SUB

    scalars = [
        ("tax_rate", "Tax rate", PCT), ("risk_free", "Risk-free rate", PCT),
        ("equity_risk_premium", "Equity risk premium", PCT),
        ("beta", "Beta", "0.00"),
        ("cost_of_debt", "Cost of debt (pre-tax)", PCT),
        ("debt_weight", "Debt weight of capital", PCT),
        (None, "WACC (calculated)", PCT),
        ("terminal_growth", "Terminal growth rate", PCT),
        ("exit_multiple", "Exit EV/EBITDA multiple", MULT),
        ("net_debt", "Net debt", NUM),
        ("shares_out", "Shares outstanding", NUM),
        ("current_price", "Current share price", CUR),
        ("wacc_step", "Sensitivity: WACC step", PCT),
        ("g_step", "Sensitivity: x-axis step", PCT),
        ("first_period_fraction", "First period (years from valuation date)", "0.00"),
        ("scenario_growth_delta", "Scenario growth shift (+/-)", PCT),
    ]
    build = spec.get("wacc_build")
    wrefs = _wacc_refs(spec) if build else {}
    for idx, (key, lab, fmt) in enumerate(scalars):
        r = 4 + idx
        label(ws, r, lab)
        if key is None:
            if build:
                put(ws, f"B{r}", f"=WACC!B{wrefs['wacc']}", fmt=fmt, fill=FILL_OUT)
                note(ws, r, bcol, "Derived on the WACC tab from the peer set, the "
                                  "regression beta, the market record and interest "
                                  "actually paid")
            else:
                put(ws, f"B{r}", "=(1-B9)*(B5+B7*B6)+B9*B8*(1-B4)", fmt=fmt,
                    fill=FILL_OUT)
                note(ws, r, bcol, "Calculated: CAPM cost of equity and after-tax "
                                  "cost of debt, weighted by the capital structure")
        elif build and key in core.BUILDUP_SUPPLIED:
            # one source of truth: these are derived on the WACC tab
            put(ws, f"B{r}", f"=WACC!B{wrefs[key]}", fmt=fmt)
            note(ws, r, bcol, "Derived on the WACC tab")
        else:
            is_est = key in est
            put(ws, f"B{r}", i.get(key, 0.0),
                font=F_ESTIMATE if is_est else F_INPUT,
                fmt=fmt, fill=FILL_EST if is_est else None)
            note(ws, r, bcol, basis.get(key, ""))
        unit = "x" if fmt in ("0.00", MULT) else ("%" if fmt == PCT
                                                  else spec["meta"]["units"])
        ws.cell(row=r, column=3, value=unit)

    hdr = I["vector_header"]
    head(ws, hdr, ["Driver"] + [f"Y{k+1}" for k in range(nf)])
    ws.cell(row=hdr, column=bcol, value="Basis / source").font = F_SUB
    vectors = [
        ("rev_growth", "Revenue growth", PCT),
        ("ebit_margin", "EBIT margin", PCT),
        ("da_pct_rev", "D&A (% of revenue)", PCT),
        ("capex_pct_rev", "Capex (% of revenue)", PCT),
        ("nwc_pct_rev", "Net working capital (% of revenue)", PCT),
    ]
    segs = spec.get("segments")
    seg_rows_map = None
    if segs:
        import build_dcf_extras as _ex
        seg_rows_map = _ex.seg_rows(len(segs))
    nh = core.n_hist(spec)
    for idx, (key, lab, fmt) in enumerate(vectors):
        r = I["rev_growth_row"] + idx
        label(ws, r, lab)
        is_est = key in est
        for t in range(nf):
            if key == "rev_growth" and segs:
                # one growth path: the one the segments imply
                put(ws, f"{col(2 + t)}{r}",
                    f"=SEGMENTS!{col(2 + nh + t)}{seg_rows_map['implied']}", fmt=fmt)
            else:
                put(ws, f"{col(2 + t)}{r}", i[key][t],
                    font=F_ESTIMATE if is_est else F_INPUT, fmt=fmt,
                    fill=FILL_EST if is_est else None)
        note(ws, r, bcol,
             "Built up on the SEGMENTS tab — growth is a consequence of the mix"
             if key == "rev_growth" and segs else basis.get(key, ""))

    comps = spec.get("comps", [])
    head(ws, I["comps_header_row"],
         ["Peer (comparables cross-check)", "EV/EBITDA", "P/E"])
    ws.cell(row=I["comps_header_row"], column=4, value="Source").font = F_SUB
    for k, peer in enumerate(comps):
        r = I["comps_start_row"] + k
        ws.cell(row=r, column=1, value=peer.get("name", "")).border = BOX
        put(ws, f"B{r}", peer.get("ev_ebitda"), font=F_INPUT, fmt=MULT)
        put(ws, f"C{r}", peer.get("pe"), font=F_INPUT, fmt=MULT)
        note(ws, r, 4, peer.get("source", ""))

    year_widths(ws, nf, label_w=36, year_w=12)
    widths(ws, {"C": 12, col(bcol): 62})
    ws.freeze_panes = "B4"
    return ws


def build_historicals(wb, spec):
    """Historicals plus a statistics block -- the reason for demanding 7-10 years."""
    ws = wb.create_sheet("HISTORICALS")
    h, src = spec["historicals"], spec["historicals"].get("sources", {})
    nh = core.n_hist(spec)
    scol = 2 + nh
    last = col(1 + nh)
    R = CELLS["HISTORICALS"]

    ws["A1"] = f"HISTORICALS — {nh} reported years, each figure with its filing"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue = as reported in the cited filing. Black = calculated "
                "from the rows above.")
    ws["A2"].font = F_NOTE

    head(ws, 3, ["Metric"] + [str(y) for y in h["years"]])
    ws.cell(row=3, column=scol, value="Source (filing, date)").font = F_SUB

    data_rows = [("revenue", "Revenue"), ("ebit", "EBIT"), ("da", "D&A"),
                 ("capex", "Capex"),
                 ("nwc_balance", "Net working capital (balance)"),
                 ("ppe_net", "Net PP&E")]
    for key, lab in data_rows:
        r = R[key]
        label(ws, r, lab)
        for t in range(nh):
            put(ws, f"{col(2 + t)}{r}", h[key][t], font=F_INPUT, fmt=NUM)
        note(ws, r, scol, src.get(key, ""))

    label(ws, R["ebit_margin"], "EBIT margin")
    label(ws, R["delta_nwc"], "Change in net working capital")
    label(ws, R["fcf"], "Unlevered FCF (derived)")
    for t in range(nh):
        c, p = col(2 + t), col(1 + t)
        put(ws, f"{c}{R['ebit_margin']}",
            f"={c}{R['ebit']}/{c}{R['revenue']}", fmt=PCT)
        if t == 0:
            put(ws, f"{c}{R['delta_nwc']}", "=NA()", fmt=NUM)
            put(ws, f"{c}{R['fcf']}", "=NA()", fmt=NUM)
        else:
            put(ws, f"{c}{R['delta_nwc']}",
                f"={c}{R['nwc_balance']}-{p}{R['nwc_balance']}", fmt=NUM)
            put(ws, f"{c}{R['fcf']}",
                f"={c}{R['ebit']}*(1-INPUTS!$B$4)+{c}{R['da']}-{c}{R['capex']}"
                f"-{c}{R['delta_nwc']}", fmt=NUM)

    # ---- statistics: the range every assumption should be judged against ---
    ws.cell(row=R["stats_header"], column=1,
            value=f"What {nh} years of history says").font = F_TITLE
    stats = [
        (R["cagr"], "Revenue CAGR over the period",
         f"=(({last}{R['revenue']}/B{R['revenue']})^(1/{nh - 1}))-1", PCT,
         "Compare this with the projected growth on INPUTS"),
        (R["margin_avg"], "EBIT margin — average",
         f"=AVERAGE(B{R['ebit_margin']}:{last}{R['ebit_margin']})", PCT, ""),
        (R["margin_min"], "EBIT margin — low",
         f"=MIN(B{R['ebit_margin']}:{last}{R['ebit_margin']})", PCT,
         "A projected margin below this has happened before; above the high has not"),
        (R["margin_max"], "EBIT margin — high",
         f"=MAX(B{R['ebit_margin']}:{last}{R['ebit_margin']})", PCT, ""),
        (R["capex_avg"], "Capex as % of revenue — average",
         f"=AVERAGE(B{R['capex']}:{last}{R['capex']})/"
         f"AVERAGE(B{R['revenue']}:{last}{R['revenue']})", PCT, ""),
        (R["down_years"], "Years revenue fell",
         f"=SUMPRODUCT(--(C{R['revenue']}:{last}{R['revenue']}<"
         f"B{R['revenue']}:{col(nh)}{R['revenue']}))", "0",
         "A model with no down years in a business that has had them is a choice"),
    ]
    for r, lab, formula, fmt, why in stats:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt, fill=FILL_BAND)
        if why:
            note(ws, r, 3, why)

    year_widths(ws, nh, label_w=36, year_w=12, extra=58)
    ws.freeze_panes = "B4"
    return ws


def build_projection(wb, spec):
    ws = wb.create_sheet("PROJECTION")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    H, P, I = CELLS["HISTORICALS"], CELLS["PROJECTION"], CELLS["INPUTS"]
    hist_last = col(1 + nh)

    ws["A1"] = "PROJECTION — every cell a formula off INPUTS"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Nothing here is typed in. Change an assumption on INPUTS and "
                "watch this move.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Metric"] + [str(y) for y in core.fcst_years(spec)])

    labels = [(P["revenue"], "Revenue"), (P["ebit"], "EBIT"),
              (P["tax"], "Tax on EBIT"), (P["nopat"], "NOPAT"),
              (P["da"], "D&A"), (P["capex"], "Capex"),
              (P["nwc_balance"], "Net working capital (balance)"),
              (P["delta_nwc"], "Change in net working capital"),
              (P["fcf"], "Unlevered free cash flow")]
    for r, lab in labels:
        label(ws, r, lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        rev_prev = (f"HISTORICALS!${hist_last}${H['revenue']}" if t == 0
                    else f"{p}{P['revenue']}")
        nwc_prev = (f"HISTORICALS!${hist_last}${H['nwc_balance']}" if t == 0
                    else f"{p}{P['nwc_balance']}")
        put(ws, f"{c}{P['revenue']}",
            f"={rev_prev}*(1+INPUTS!{c}${I['rev_growth_row']})", fmt=NUM)
        put(ws, f"{c}{P['ebit']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['ebit_margin_row']}", fmt=NUM)
        put(ws, f"{c}{P['tax']}", f"={c}{P['ebit']}*INPUTS!$B$4", fmt=NUM)
        put(ws, f"{c}{P['nopat']}", f"={c}{P['ebit']}-{c}{P['tax']}", fmt=NUM)
        put(ws, f"{c}{P['da']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['da_pct_rev_row']}", fmt=NUM)
        put(ws, f"{c}{P['capex']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['capex_pct_rev_row']}", fmt=NUM)
        put(ws, f"{c}{P['nwc_balance']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['nwc_pct_rev_row']}", fmt=NUM)
        put(ws, f"{c}{P['delta_nwc']}",
            f"={c}{P['nwc_balance']}-{nwc_prev}", fmt=NUM)
        put(ws, f"{c}{P['fcf']}",
            f"={c}{P['nopat']}+{c}{P['da']}-{c}{P['capex']}-{c}{P['delta_nwc']}",
            fmt=NUM, fill=FILL_OUT)

    year_widths(ws, nf, label_w=36)
    ws.freeze_panes = "B4"
    return ws


def build_valuation(wb, spec):
    ws = wb.create_sheet("VALUATION")
    nf = core.n_fcst(spec)
    V, P, I = CELLS["VALUATION"], CELLS["PROJECTION"], CELLS["INPUTS"]
    last = col(1 + nf)

    ws["A1"] = "VALUATION — DCF with an exit-multiple and comparables cross-check"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Outputs a value range under stated assumptions. It does not say "
                "whether to buy or sell.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Metric"] + [f"Y{k+1}" for k in range(nf)])

    for r, lab in ((V["fcf"], "Unlevered free cash flow"),
                   (V["period"], "Discount period"),
                   (V["discount_factor"], "Discount factor"),
                   (V["pv_fcf"], "PV of free cash flow")):
        label(ws, r, lab)
    for t in range(nf):
        c = col(2 + t)
        put(ws, f"{c}{V['fcf']}", f"=PROJECTION!{c}{P['fcf']}", fmt=NUM)
        put(ws, f"{c}{V['period']}",
            f"=INPUTS!$B$18+COLUMN({c}{V['period']})-2", fmt="0.00")
        put(ws, f"{c}{V['discount_factor']}",
            f"=1/(1+INPUTS!$B$10)^{c}{V['period']}", fmt="0.0000")
        put(ws, f"{c}{V['pv_fcf']}",
            f"={c}{V['fcf']}*{c}{V['discount_factor']}", fmt=NUM)

    bridge = [
        (V["pv_explicit"], "PV of explicit-period FCF",
         f"=SUM(B{V['pv_fcf']}:{last}{V['pv_fcf']})", NUM, None),
        (V["terminal_value"], "Terminal value (perpetuity growth)",
         f"={last}{V['fcf']}*(1+INPUTS!$B$11)/(INPUTS!$B$10-INPUTS!$B$11)", NUM, None),
        (V["pv_terminal"], "PV of terminal value",
         f"={V['terminal_value']}*{last}{V['discount_factor']}", NUM, None),
        (V["enterprise_value"], "Enterprise value",
         f"={V['pv_explicit']}+{V['pv_terminal']}", NUM, None),
        (V["net_debt"], "Less: net debt", "=INPUTS!$B$13", NUM, None),
        (V["equity_value"], "Equity value",
         f"={V['enterprise_value']}-{V['net_debt']}", NUM, None),
        (V["shares"], "Shares outstanding", "=INPUTS!$B$14", NUM, None),
        (V["value_per_share"], "Value per share (DCF, perpetuity)",
         f"={V['equity_value']}/{V['shares']}", CUR, FILL_OUT),
        (V["current_price"], "Current share price", "=INPUTS!$B$15", CUR, None),
        (V["implied_vs_price"], "Model value vs current price (arithmetic only)",
         f"={V['value_per_share']}/{V['current_price']}-1", PCT, None),
    ]
    for addr, lab, formula, fmt, fill in bridge:
        label(ws, int(addr[1:]), lab)
        put(ws, addr, formula, fmt=fmt, fill=fill)

    ws["A20"] = "Cross-check 1 — exit multiple terminal value"
    ws["A20"].font = F_SUB
    for addr, lab, formula, fmt in [
        (V["terminal_ebitda"], "Terminal-year EBITDA",
         f"=PROJECTION!{last}{P['ebit']}+PROJECTION!{last}{P['da']}", NUM),
        (V["tv_exit"], "Terminal value (exit multiple)",
         f"={V['terminal_ebitda']}*INPUTS!$B$12", NUM),
        (V["pv_tv_exit"], "PV of terminal value (exit)",
         f"={V['tv_exit']}*{last}{V['discount_factor']}", NUM),
        (V["ev_exit"], "Enterprise value (exit method)",
         f"={V['pv_explicit']}+{V['pv_tv_exit']}", NUM),
        (V["value_per_share_exit"], "Value per share (exit method)",
         f"=({V['ev_exit']}-INPUTS!$B$13)/INPUTS!$B$14", CUR),
    ]:
        label(ws, int(addr[1:]), lab)
        put(ws, addr, formula, fmt=fmt)

    ws["A27"] = "Cross-check 2 — comparables"
    ws["A27"].font = F_SUB
    n_comps = max(len(spec.get("comps", [])), 1)
    clast = CELLS["INPUTS"]["comps_start_row"] + n_comps - 1
    for addr, lab, formula, fmt in [
        (V["peer_median_ev_ebitda"], "Peer median EV/EBITDA",
         f"=MEDIAN(INPUTS!B{I['comps_start_row']}:B{clast})", MULT),
        (V["ntm_ebitda"], "Y1 EBITDA",
         f"=PROJECTION!B{P['ebit']}+PROJECTION!B{P['da']}", NUM),
        (V["implied_ev_comps"], "Implied enterprise value (comps)",
         f"={V['peer_median_ev_ebitda']}*{V['ntm_ebitda']}", NUM),
        (V["value_per_share_comps"], "Implied value per share (comps)",
         f"=({V['implied_ev_comps']}-INPUTS!$B$13)/INPUTS!$B$14", CUR),
    ]:
        label(ws, int(addr[1:]), lab)
        put(ws, addr, formula, fmt=fmt)

    year_widths(ws, nf, label_w=42)
    return ws


AXIS_CONFIG = {
    "terminal_growth": {"label": "Terminal growth rate", "fmt": PCT,
                        "base": "INPUTS!$B$11", "driver": None},
    "rev_growth": {"label": "Revenue growth (parallel shift)", "fmt": PCT,
                   "base": "INPUTS!$B$22", "driver": "rev_growth"},
    "ebit_margin": {"label": "EBIT margin (parallel shift)", "fmt": PCT,
                    "base": "INPUTS!$B$23", "driver": "ebit_margin"},
}


def build_sensitivity(wb, spec, x_axis="terminal_growth"):
    """Two-way grid. Y is always WACC; X is configurable.

    For an operating X-axis the grid cannot reuse the base projection, so the
    sheet carries a helper block per column: a full formula-driven mini
    projection with the shift applied. Nothing in the grid is hardcoded.
    """
    if x_axis not in AXIS_CONFIG:
        raise ValueError(f"unsupported sensitivity axis: {x_axis}")
    cfg = AXIS_CONFIG[x_axis]
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    V, P, H, I, S = (CELLS["VALUATION"], CELLS["PROJECTION"],
                     CELLS["HISTORICALS"], CELLS["INPUTS"], CELLS["SENSITIVITY"])
    hist_last, flast = col(1 + nh), col(1 + nf)

    ws = wb.create_sheet("SENSITIVITY")
    ws["A1"] = f"SENSITIVITY — value per share: WACC vs {cfg['label']}"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Axis values are driven off INPUTS by the step cells (B16/B17). "
                "Change a step and the grid re-spaces itself.")
    ws["A2"].font = F_NOTE
    ws.cell(row=S["x_row"] - 1, column=2,
            value=f"WACC ↓ / {cfg['label']} →").font = F_SUB

    for k in range(GRID):
        c = col(S["grid_first_col"] + k)
        put(ws, f"{c}{S['x_row']}",
            f"={cfg['base']}+(COLUMN({c}{S['x_row']})-"
            f"{S['grid_first_col'] + GRID // 2})*INPUTS!$B$17",
            font=F_SUB, fmt=cfg["fmt"], fill=FILL_BAND)
    for j in range(GRID):
        r = S["y_start"] + j
        put(ws, f"B{r}",
            f"=INPUTS!$B$10+(ROW(B{r})-{S['y_start'] + GRID // 2})*INPUTS!$B$16",
            font=F_SUB, fmt=PCT, fill=FILL_BAND)

    hstart = S["helper_start"]
    ws.cell(row=hstart, column=1,
            value="Helper blocks — one formula-driven projection per grid column "
                  "(that column's shift applied to the INPUTS driver).").font = F_NOTE

    block_rows = {}
    for k in range(GRID):
        xcol = col(S["grid_first_col"] + k)
        base = hstart + 2 + k * (nf + 6)
        block_rows[k] = base
        ws.cell(row=base, column=1,
                value=f"Column {k + 1}: {cfg['label']} = ").font = F_SUB
        put(ws, f"B{base}", f"={xcol}${S['x_row']}", font=F_SUB, fmt=cfg["fmt"])
        shift = ("" if cfg["driver"] is None
                 else f"+({xcol}${S['x_row']}-{cfg['base']})")
        for n, lab in enumerate(["Revenue", "EBIT", "NOPAT", "D&A", "Capex",
                                 "NWC balance", "Change in NWC",
                                 "Unlevered FCF"]):
            ws.cell(row=base + 1 + n, column=1, value=lab).font = F_NOTE

        r_rev, r_ebit, r_nopat = base + 1, base + 2, base + 3
        r_da, r_capex, r_nwc = base + 4, base + 5, base + 6
        r_dnwc, r_fcf = base + 7, base + 8
        for t in range(nf):
            c, p = col(2 + t), col(1 + t)
            g_shift = shift if x_axis == "rev_growth" else ""
            m_shift = shift if x_axis == "ebit_margin" else ""
            prev_rev = (f"HISTORICALS!${hist_last}${H['revenue']}" if t == 0
                        else f"{p}{r_rev}")
            prev_nwc = (f"HISTORICALS!${hist_last}${H['nwc_balance']}" if t == 0
                        else f"{p}{r_nwc}")
            put(ws, f"{c}{r_rev}",
                f"={prev_rev}*(1+INPUTS!{c}${I['rev_growth_row']}{g_shift})", fmt=NUM)
            put(ws, f"{c}{r_ebit}",
                f"={c}{r_rev}*(INPUTS!{c}${I['ebit_margin_row']}{m_shift})", fmt=NUM)
            put(ws, f"{c}{r_nopat}", f"={c}{r_ebit}*(1-INPUTS!$B$4)", fmt=NUM)
            put(ws, f"{c}{r_da}",
                f"={c}{r_rev}*INPUTS!{c}${I['da_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{r_capex}",
                f"={c}{r_rev}*INPUTS!{c}${I['capex_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{r_nwc}",
                f"={c}{r_rev}*INPUTS!{c}${I['nwc_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{r_dnwc}", f"={c}{r_nwc}-{prev_nwc}", fmt=NUM)
            put(ws, f"{c}{r_fcf}",
                f"={c}{r_nopat}+{c}{r_da}-{c}{r_capex}-{c}{r_dnwc}", fmt=NUM)

    for k in range(GRID):
        xcol = col(S["grid_first_col"] + k)
        fcf_row = block_rows[k] + 8
        g_ref = (f"{xcol}${S['x_row']}" if x_axis == "terminal_growth"
                 else "INPUTS!$B$11")
        for j in range(GRID):
            r = S["y_start"] + j
            w = f"$B{r}"
            pv = (f"SUMPRODUCT($B{fcf_row}:${flast}{fcf_row},"
                  f"1/(1+{w})^VALUATION!$B${V['period']}:${flast}${V['period']})")
            tv = (f"${flast}{fcf_row}*(1+{g_ref})/({w}-{g_ref})"
                  f"/(1+{w})^VALUATION!${flast}${V['period']}")
            put(ws, f"{xcol}{r}",
                f"=IF({w}<={g_ref},NA(),({pv}+{tv}-INPUTS!$B$13)/INPUTS!$B$14)",
                fmt=CUR,
                fill=FILL_OUT if (j == GRID // 2 and k == GRID // 2) else None)

    ws.cell(row=S["y_start"] + GRID + 1, column=2,
            value="Shaded centre cell = base case; it must equal VALUATION value "
                  "per share (CHECKS verifies this).").font = F_NOTE
    year_widths(ws, GRID + 1, label_w=26, year_w=14)
    return ws


def build_dcf_checks(wb, spec, scen_map=None, rel_map=None):
    import build_dcf_extras as _ex
    seg_check = (_ex.seg_rows(len(spec["segments"]))["check"]
                 if spec.get("segments") else None)
    nf = core.n_fcst(spec)
    V, P, S = CELLS["VALUATION"], CELLS["PROJECTION"], CELLS["SENSITIVITY"]
    flast = col(1 + nf)
    cc = col(S["grid_first_col"] + GRID // 2)
    cr = S["y_start"] + GRID // 2
    return checks_sheet(wb, [
        ("Enterprise value ties to its components",
         f'=IF(ABS(VALUATION!{V["enterprise_value"]}-(VALUATION!{V["pv_explicit"]}'
         f'+VALUATION!{V["pv_terminal"]}))<0.01,"PASS","FAIL")',
         "EV = PV of explicit FCF + PV of terminal value"),
        ("Equity bridge ties",
         f'=IF(ABS(VALUATION!{V["equity_value"]}-(VALUATION!{V["enterprise_value"]}'
         f'-VALUATION!{V["net_debt"]}))<0.01,"PASS","FAIL")',
         "Equity value = EV less net debt"),
        ("Terminal growth below WACC",
         '=IF(INPUTS!$B$11<INPUTS!$B$10,"PASS","FAIL")',
         "The perpetuity formula converges"),
        ("Y1 free cash flow identity",
         f'=IF(ABS(PROJECTION!B{P["fcf"]}-(PROJECTION!B{P["nopat"]}'
         f'+PROJECTION!B{P["da"]}-PROJECTION!B{P["capex"]}'
         f'-PROJECTION!B{P["delta_nwc"]}))<0.01,"PASS","FAIL")',
         "FCF = NOPAT + D&A - capex - change in working capital"),
        ("Final-year free cash flow identity",
         f'=IF(ABS(PROJECTION!{flast}{P["fcf"]}-(PROJECTION!{flast}{P["nopat"]}'
         f'+PROJECTION!{flast}{P["da"]}-PROJECTION!{flast}{P["capex"]}'
         f'-PROJECTION!{flast}{P["delta_nwc"]}))<0.01,"PASS","FAIL")',
         "The identity holds at the far end of the projection too"),
        ("Revenue growth matches the input",
         f'=IF(ABS(PROJECTION!C{P["revenue"]}/PROJECTION!B{P["revenue"]}-1'
         f'-INPUTS!C$22)<0.0001,"PASS","FAIL")',
         "The projection is actually driven by the growth assumption"),
        ("Discount factors decline",
         f'=IF(VALUATION!{flast}{V["discount_factor"]}'
         f'<VALUATION!B{V["discount_factor"]},"PASS","FAIL")',
         "Discounting runs in the right direction"),
        ("Sensitivity centre equals base case",
         f'=IF(ABS(SENSITIVITY!{cc}{cr}-VALUATION!{V["value_per_share"]})'
         f'<0.01,"PASS","FAIL")',
         "The grid is wired to the same model, not a separate one"),
        ("Share count positive",
         '=IF(INPUTS!$B$14>0,"PASS","FAIL")',
         "No divide-by-zero in per-share figures"),
        ("Projected margin within historical range",
         f'=IF(AND(INPUTS!B23<=HISTORICALS!$B${CELLS["HISTORICALS"]["margin_max"]},'
         f'INPUTS!B23>=HISTORICALS!$B${CELLS["HISTORICALS"]["margin_min"]}),'
         f'"PASS","REVIEW")',
         "Year-1 margin sits inside what the company has actually achieved; "
         "REVIEW means justify it, not hide it"),
        ("Segments reconcile to reported revenue",
         (f'=IF(SUMPRODUCT(MAX(ABS(SEGMENTS!B{seg_check}:{col(1 + core.n_hist(spec))}'
          f'{seg_check})))<0.01,"PASS","FAIL")') if seg_check else '="PASS"',
         "Every reported year's segments add up to the revenue in the accounts"),
        ("Scenario base case equals the headline valuation",
         (f'=IF(ABS(SCENARIOS!C{scen_map["base"]}-VALUATION!{V["value_per_share"]})'
          f'<0.01,"PASS","FAIL")') if scen_map else '="PASS"',
         "The bear/bull cases are re-runs of the same model, not haircuts"),
        ("Scenarios ordered bear < base < bull",
         (f'=IF(AND(SCENARIOS!C{scen_map["bear"]}<SCENARIOS!C{scen_map["base"]},'
          f'SCENARIOS!C{scen_map["base"]}<SCENARIOS!C{scen_map["bull"]}),'
          f'"PASS","FAIL")') if scen_map else '="PASS"',
         "Higher growth produces a higher value, or something is wired backwards"),
        ("Assumed growth at or below intrinsic growth",
         f'=IF(INTRINSIC_GROWTH!B{20}<=INTRINSIC_GROWTH!B{20},'
         f'IF(INPUTS!B22<=INTRINSIC_GROWTH!B20,"PASS","REVIEW"),"REVIEW")',
         "Growth above RR x ROIC has to come from somewhere; REVIEW means say "
         "where, not hide it"),
        ("Exit cross-check within 50% of DCF",
         f'=IF(ABS(VALUATION!{V["value_per_share_exit"]}'
         f'/VALUATION!{V["value_per_share"]}-1)<0.5,"PASS","REVIEW")',
         "Two terminal methods broadly agree; REVIEW means explain the gap"),
    ])


def build_dcf(spec, out_path, x_axis="terminal_growth"):
    import build_dcf_extras as ex

    wb = Workbook()
    wb.remove(wb.active)
    stub = spec["inputs"].get("first_period_fraction", 1.0)
    notes = [
        ("History", f"{core.n_hist(spec)} reported years; see the statistics "
                    "block on HISTORICALS for the range each assumption sits in"),
        ("Forecast", f"{core.n_fcst(spec)} years, then a terminal value"),
        ("Discounting", f"First period is {stub:.2f} years from the valuation "
                        "date; 1.00 is year-end, 0.50 the mid-year convention"),
    ]
    if spec.get("wacc_build"):
        notes.append(("WACC", "Derived on the WACC tab — peer betas unlevered and "
                              "re-levered, a Blume-adjusted regression beta, an "
                              "equity risk premium from the market record, and the "
                              "cost of debt from interest actually paid"))
    if spec.get("segments"):
        notes.append(("Top line", "Built up on the SEGMENTS tab; the group growth "
                                  "rate on INPUTS is derived from the mix"))
    if (spec.get("wacc_build") or {}).get("regression_beta", {}).get("prices"):
        notes.append(("Beta", "Regressed from the price series on "
                              "BETA_REGRESSION — slope, alpha and R² computed "
                              "in the sheet"))
    if spec.get("relative_peers"):
        notes.append(("Cross-checks", "RELATIVE_PRICING carries the full peer "
                                      "distribution; FOOTBALL_FIELD puts every "
                                      "method's range on one scale"))
    readme_sheet(wb, spec, "DCF valuation model", notes)

    build_inputs(wb, spec)
    build_historicals(wb, spec)
    if spec.get("segments"):
        ex.build_segments(wb, spec, CELLS["HISTORICALS"])
    reg_ref = None
    if (spec.get("wacc_build") or {}).get("regression_beta", {}).get("prices"):
        _, BS = ex.build_beta_regression(wb, spec)
        reg_ref = f"BETA_REGRESSION!H{BS['slope']}"
    if spec.get("wacc_build"):
        ex.build_wacc(wb, spec, reg_ref)
    build_projection(wb, spec)
    build_valuation(wb, spec)
    if spec.get("wacc_build"):
        tax_ref = f"WACC!$B${_wacc_refs(spec) and ex.erp_rows(ex.wacc_rows(len(spec['wacc_build'].get('peers', []))), len(spec['wacc_build'].get('market_returns', {}).get('years', [])))['tax_rate']}"
    else:
        tax_ref = "INPUTS!$B$4"
    ex.build_intrinsic(wb, spec, CELLS["HISTORICALS"], tax_ref)
    rel_map = None
    if spec.get("relative_peers"):
        _, rel_map = ex.build_relative(wb, spec)
    build_sensitivity(wb, spec, x_axis)
    _, scen_map = ex.build_scenarios(wb, spec, CELLS)
    if rel_map:
        ex.build_football(wb, spec, rel_map, scen_map, CELLS)
    build_dcf_checks(wb, spec, scen_map, rel_map)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------
# dispatch
# --------------------------------------------------------------------------

BUILDERS = {"dcf": build_dcf}


def load_for_type(path, model_type):
    import json
    with open(path, "r", encoding="utf-8") as fh:
        spec = json.load(fh)
    if model_type == "dcf":
        core.validate_spec(spec)
    elif model_type == "statements":
        import stmt_core
        stmt_core.validate(spec)
    elif model_type == "lbo":
        import lbo_core
        lbo_core.validate(spec)
    elif model_type == "ipo":
        import ipo_core
        ipo_core.validate(spec)
    return spec


def register_builders():
    try:
        import build_statements
        BUILDERS["statements"] = build_statements.build
    except ImportError:
        pass
    try:
        import build_lbo
        BUILDERS["lbo"] = build_lbo.build
    except ImportError:
        pass
    try:
        import build_ipo
        BUILDERS["ipo"] = build_ipo.build
    except ImportError:
        pass


def main():
    register_builders()
    ap = argparse.ArgumentParser(description="Financopedia model builder")
    ap.add_argument("spec")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--type", default="dcf",
                    choices=["dcf", "statements", "lbo", "ipo"])
    ap.add_argument("--x-axis", default="terminal_growth", choices=sorted(AXIS_CONFIG))
    ap.add_argument("--scan", action="store_true")
    ap.add_argument("--breakeven", default=None)
    ap.add_argument("--history", action="store_true",
                    help="print what the historical record says")
    ap.add_argument("--move", type=float, default=0.30)
    args = ap.parse_args()

    spec = load_for_type(args.spec, args.type)

    if args.history:
        s = core.history_stats(spec)
        print(f"\n{spec['meta']['name']} — {s['years']} years of history")
        print(f"  revenue CAGR          {s['revenue_cagr']:>8.1%}")
        print(f"  revenue growth  min {s['revenue_growth']['min']:>7.1%}  "
              f"avg {s['revenue_growth']['avg']:>7.1%}  "
              f"max {s['revenue_growth']['max']:>7.1%}")
        print(f"  EBIT margin     min {s['ebit_margin']['min']:>7.1%}  "
              f"avg {s['ebit_margin']['avg']:>7.1%}  "
              f"max {s['ebit_margin']['max']:>7.1%}")
        print(f"  capex % revenue min {s['capex_pct_rev']['min']:>7.1%}  "
              f"avg {s['capex_pct_rev']['avg']:>7.1%}  "
              f"max {s['capex_pct_rev']['max']:>7.1%}")
        print(f"  years revenue fell    {s['down_years']:>8}")
        flags = core.assumption_flags(spec)
        if flags:
            print("\n  assumptions outside the historical range:")
            for f in flags:
                print(f"    - {f}")
        else:
            print("\n  every forecast driver sits inside the historical range")
        return 0

    if args.scan:
        rows = core.scan(spec)
        print(f"{'input':<22}{'base':>10}{'-10%':>12}{'+10%':>12}"
              f"{'swing':>12}{'% of base':>12}")
        print("-" * 80)
        for r in rows:
            print(f"{r['input']:<22}{r['base']:>10.2f}{r['down']:>12.2f}"
                  f"{r['up']:>12.2f}{r['swing']:>12.2f}{r['swing_pct']:>11.1%}")
        print("\nTop two by swing:", ", ".join(r["input"] for r in rows[:2]))
        return 0

    if args.breakeven:
        for key in [k.strip() for k in args.breakeven.split(",")]:
            res = core.breakeven(spec, key, args.move)
            print(f"\n{key}: base value/share {res['base_value']:.2f}")
            for d in ("up", "down"):
                rel = res.get(d)
                print(f"  {d:<5} " + ("not solvable in range" if rel is None
                                       else f"requires {rel:+.1%} change in {key}"))
        return 0

    if args.type not in BUILDERS:
        print(f"builder for --type {args.type} is not installed", file=sys.stderr)
        return 2

    out = args.out or f"{spec['meta']['ticker']}-{args.type}.xlsx"
    if args.type == "dcf":
        BUILDERS["dcf"](spec, out, args.x_axis)
        v = core.run(spec)
        print(f"built {out}")
        print(f"  {core.n_hist(spec)}y history · {core.n_fcst(spec)}y forecast · "
              f"WACC {v.wacc:.2%} · EV {v.enterprise_value:,.0f} · "
              f"value/share {v.value_per_share:,.2f}")
    else:
        BUILDERS[args.type](spec, out)
        print(f"built {out}")
    print("  run audit_model.py before delivering this file")
    return 0


if __name__ == "__main__":
    sys.exit(main())
