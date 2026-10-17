#!/usr/bin/env python3
"""Three-statement workbook: linked income statement, balance sheet and cash flow.

The balance sheet balances by construction, which is why CHECKS tests it: if
it ever stops balancing, something upstream has broken.
"""

from __future__ import annotations

from openpyxl import Workbook

import dcf_core as core
import stmt_core as st
from xl_style import (CUR, DAYS, F_ESTIMATE, F_INPUT, F_NOTE, F_SUB, F_TITLE,
                      FILL_BAND, FILL_EST, FILL_OUT, MULT, NUM, PCT,
                      checks_sheet, col, head, label, note, put, readme_sheet,
                      widths, year_widths)

STMT = {
    "INPUTS": {
        "tax_rate": "B4", "interest_rate_debt": "B5",
        "interest_rate_revolver": "B6", "interest_rate_cash": "B7",
        "min_cash": "B8", "dividend_payout": "B9", "opening_revolver": "B10",
        "rev_growth": 13, "gross_margin": 14, "sga_pct_rev": 15,
        "da_pct_rev": 16, "capex_pct_rev": 17, "dso": 18, "dio": 19, "dpo": 20,
        "other_ca_pct_rev": 21, "accrued_pct_rev": 22, "debt_repayment": 23,
    },
    "HISTORICALS": {
        "revenue": 4, "cogs": 5, "gross_profit": 6, "sga": 7, "ebitda": 8,
        "da": 9, "ebit": 10, "capex": 11, "net_income": 36,
        "interest_expense": 37, "cfo": 38,
        "cash": 13, "ar": 14, "inventory": 15, "other_ca": 16, "ppe_net": 17,
        "other_assets": 18, "total_assets": 19,
        "ap": 21, "accrued": 22, "other_cl": 23, "debt": 24,
        "other_lt_liab": 25, "equity": 26, "total_le": 27, "balance": 28,
        "nwc": 29,
        "cagr": 41, "gm_avg": 42, "em_avg": 43, "em_min": 44,
    },
    "INCOME": {
        "revenue": 4, "cogs": 5, "gross_profit": 6, "sga": 7, "ebitda": 8,
        "da": 9, "ebit": 10, "interest_expense": 11, "interest_income": 12,
        "pbt": 13, "tax": 14, "net_income": 15, "dividends": 16, "retained": 17,
    },
    "BALANCE": {
        "cash": 4, "ar": 5, "inventory": 6, "other_ca": 7, "ppe_net": 8,
        "other_assets": 9, "total_assets": 10,
        "ap": 12, "accrued": 13, "other_cl": 14, "debt": 15, "revolver": 16,
        "other_lt_liab": 17, "equity": 18, "total_le": 19, "balance": 21,
        "nwc": 23, "net_debt": 24,
    },
    "CASHFLOW": {
        "net_income": 4, "da": 5, "delta_wc": 6, "cfo": 7, "capex": 8,
        "cfi": 9, "debt_repayment": 10, "dividends": 11, "opening_cash": 12,
        "cash_before_revolver": 13, "revolver_draw": 14, "revolver_repay": 15,
        "cff": 16, "net_change": 17, "closing_cash": 18,
        "unlevered_fcf": 20, "levered_fcf": 21,
    },
    "RATIOS": {
        "revenue_growth": 4, "ebitda_growth": 5, "gross_margin": 6,
        "ebitda_margin": 7, "ebit_margin": 8, "net_margin": 9, "capex_pct": 10,
        "dso": 11, "dio": 12, "dpo": 13, "ccc": 14, "fixed_asset_turnover": 15,
        "net_debt_ebitda": 16, "interest_cover": 17, "roce": 18, "roe": 19,
        "retention": 20, "sustainable_growth": 21, "cfo_sales": 22,
        "cfo_assets": 23, "cfo_debt": 24, "fcf_conversion": 25,
    },
}

D = 365


def build_inputs(wb, spec):
    ws = wb.create_sheet("INPUTS")
    i = spec["inputs"]
    basis, est = i.get("basis", {}), set(i.get("estimates", []))
    nf = core.n_fcst(spec)
    bcol = max(2 + nf, 5)
    I = STMT["INPUTS"]

    ws["A1"] = "INPUTS — every assumption behind the three statements"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue = input. Amber italic = estimated. Interest is charged on "
                "opening balances, so the model needs no iterative calculation.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Assumption", "Value", "Unit"])
    ws.cell(row=3, column=bcol, value="Basis / source").font = F_SUB

    scalars = [
        ("tax_rate", "Tax rate", PCT, "%"),
        ("interest_rate_debt", "Interest rate — term debt", PCT, "%"),
        ("interest_rate_revolver", "Interest rate — revolver", PCT, "%"),
        ("interest_rate_cash", "Interest earned on cash", PCT, "%"),
        ("min_cash", "Minimum cash balance", NUM, spec["meta"]["units"]),
        ("dividend_payout", "Dividend payout ratio", PCT, "%"),
        ("opening_revolver", "Opening revolver balance", NUM, spec["meta"]["units"]),
    ]
    for key, lab, fmt, unit in scalars:
        r = int(I[key][1:])
        label(ws, r, lab)
        is_est = key in est
        put(ws, I[key], i[key], font=F_ESTIMATE if is_est else F_INPUT, fmt=fmt,
            fill=FILL_EST if is_est else None)
        ws.cell(row=r, column=3, value=unit)
        note(ws, r, bcol, basis.get(key, ""))

    head(ws, 12, ["Forecast driver"] + [f"Y{k+1}" for k in range(nf)])
    ws.cell(row=12, column=bcol, value="Basis / source").font = F_SUB
    vectors = [
        ("rev_growth", "Revenue growth", PCT),
        ("gross_margin", "Gross margin", PCT),
        ("sga_pct_rev", "SG&A (% of revenue)", PCT),
        ("da_pct_rev", "D&A (% of revenue)", PCT),
        ("capex_pct_rev", "Capex (% of revenue)", PCT),
        ("dso", "Days sales outstanding", DAYS),
        ("dio", "Days inventory outstanding", DAYS),
        ("dpo", "Days payables outstanding", DAYS),
        ("other_ca_pct_rev", "Other current assets (% of revenue)", PCT),
        ("accrued_pct_rev", "Accrued liabilities (% of revenue)", PCT),
        ("debt_repayment", "Scheduled term debt repayment", NUM),
    ]
    for key, lab, fmt in vectors:
        r = I[key]
        label(ws, r, lab)
        is_est = key in est
        for t in range(nf):
            put(ws, f"{col(2 + t)}{r}", i[key][t],
                font=F_ESTIMATE if is_est else F_INPUT, fmt=fmt,
                fill=FILL_EST if is_est else None)
        note(ws, r, bcol, basis.get(key, ""))

    year_widths(ws, nf, label_w=38, year_w=12)
    widths(ws, {"C": 12, col(bcol): 60})
    return ws


def build_historicals(wb, spec):
    ws = wb.create_sheet("HISTORICALS")
    h, src = spec["historicals"], spec["historicals"].get("sources", {})
    nh = core.n_hist(spec)
    scol, last = 2 + nh, col(1 + nh)
    H = STMT["HISTORICALS"]

    ws["A1"] = f"HISTORICALS — {nh} reported years of income statement and balance sheet"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue = as reported. The balance row proves each year's reported "
                "balance sheet actually balances.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Line item"] + [str(y) for y in h["years"]])
    ws.cell(row=3, column=scol, value="Source (filing, date)").font = F_SUB

    data = [("net_income", "Net income"),
            ("interest_expense", "Interest expense"),
            ("cfo", "Cash from operating activities"),
            ("revenue", "Revenue"), ("cogs", "Cost of goods sold"),
            ("sga", "SG&A"), ("da", "D&A"), ("capex", "Capex"),
            ("cash", "Cash and equivalents"), ("ar", "Accounts receivable"),
            ("inventory", "Inventory"), ("other_ca", "Other current assets"),
            ("ppe_net", "Net PP&E"), ("other_assets", "Other assets"),
            ("ap", "Accounts payable"), ("accrued", "Accrued liabilities"),
            ("other_cl", "Other current liabilities"), ("debt", "Total debt"),
            ("other_lt_liab", "Other long-term liabilities"),
            ("equity", "Shareholders' equity")]
    for key, lab in data:
        r = H[key]
        label(ws, r, lab)
        for t in range(nh):
            put(ws, f"{col(2 + t)}{r}", h[key][t], font=F_INPUT, fmt=NUM)
        note(ws, r, scol, src.get(key, ""))

    for r, lab in ((H["gross_profit"], "Gross profit"),
                   (H["ebitda"], "EBITDA"), (H["ebit"], "EBIT"),
                   (H["total_assets"], "Total assets"),
                   (H["total_le"], "Total liabilities and equity"),
                   (H["balance"], "Balance check (assets − L&E)"),
                   (H["nwc"], "Net working capital (memo)")):
        label(ws, r, lab)

    for t in range(nh):
        c = col(2 + t)
        put(ws, f"{c}{H['gross_profit']}", f"={c}{H['revenue']}-{c}{H['cogs']}", fmt=NUM)
        put(ws, f"{c}{H['ebitda']}",
            f"={c}{H['gross_profit']}-{c}{H['sga']}", fmt=NUM)
        put(ws, f"{c}{H['ebit']}", f"={c}{H['ebitda']}-{c}{H['da']}", fmt=NUM)
        put(ws, f"{c}{H['total_assets']}",
            f"=SUM({c}{H['cash']}:{c}{H['other_assets']})", fmt=NUM, fill=FILL_BAND)
        put(ws, f"{c}{H['total_le']}",
            f"=SUM({c}{H['ap']}:{c}{H['equity']})", fmt=NUM, fill=FILL_BAND)
        put(ws, f"{c}{H['balance']}",
            f"={c}{H['total_assets']}-{c}{H['total_le']}", fmt=NUM)
        put(ws, f"{c}{H['nwc']}",
            f"={c}{H['ar']}+{c}{H['inventory']}+{c}{H['other_ca']}"
            f"-{c}{H['ap']}-{c}{H['accrued']}", fmt=NUM)

    ws.cell(row=35, column=1,
            value="Income statement and cash flow, reported").font = F_TITLE
    ws.cell(row=40, column=1,
            value=f"What {nh} years of history says").font = F_TITLE
    for r, lab, formula, fmt, why in [
        (H["cagr"], "Revenue CAGR over the period",
         f"=(({last}{H['revenue']}/B{H['revenue']})^(1/{nh - 1}))-1", PCT,
         "Compare with the projected growth on INPUTS"),
        (H["gm_avg"], "Gross margin — average",
         f"=AVERAGE(B{H['gross_profit']}:{last}{H['gross_profit']})"
         f"/AVERAGE(B{H['revenue']}:{last}{H['revenue']})", PCT, ""),
        (H["em_avg"], "EBITDA margin — average",
         f"=AVERAGE(B{H['ebitda']}:{last}{H['ebitda']})"
         f"/AVERAGE(B{H['revenue']}:{last}{H['revenue']})", PCT, ""),
        (H["em_min"], "EBITDA margin — worst year",
         f"=MIN(B{H['ebitda']}:{last}{H['ebitda']}/B{H['revenue']}:{last}{H['revenue']})",
         PCT, "The trough the forecast should be tested against"),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt, fill=FILL_BAND)
        if why:
            note(ws, r, 3, why)

    year_widths(ws, nh, label_w=36, year_w=12, extra=52)
    ws.freeze_panes = "B4"
    return ws


def build_income(wb, spec):
    ws = wb.create_sheet("INCOME")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    HL = col(1 + nh)
    I, H, B, N = STMT["INPUTS"], STMT["HISTORICALS"], STMT["BALANCE"], STMT["INCOME"]

    ws["A1"] = "INCOME STATEMENT — forecast"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Interest is charged on opening balances, which is why no cell "
                "here refers to this year's closing debt.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Line item"] + [str(y) for y in core.fcst_years(spec)])

    for key, lab in [("revenue", "Revenue"), ("cogs", "Cost of goods sold"),
                     ("gross_profit", "Gross profit"), ("sga", "SG&A"),
                     ("ebitda", "EBITDA"), ("da", "D&A"), ("ebit", "EBIT"),
                     ("interest_expense", "Interest expense"),
                     ("interest_income", "Interest income"),
                     ("pbt", "Profit before tax"), ("tax", "Tax"),
                     ("net_income", "Net income"),
                     ("dividends", "Dividends declared"),
                     ("retained", "Retained earnings")]:
        label(ws, N[key], lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        prev_rev = (f"HISTORICALS!${HL}${H['revenue']}" if t == 0
                    else f"{p}{N['revenue']}")
        open_debt = (f"HISTORICALS!${HL}${H['debt']}" if t == 0
                     else f"BALANCE!{p}{B['debt']}")
        open_rev = ("INPUTS!$B$10" if t == 0 else f"BALANCE!{p}{B['revolver']}")
        open_cash = (f"HISTORICALS!${HL}${H['cash']}" if t == 0
                     else f"BALANCE!{p}{B['cash']}")

        put(ws, f"{c}{N['revenue']}",
            f"={prev_rev}*(1+INPUTS!{c}${I['rev_growth']})", fmt=NUM)
        put(ws, f"{c}{N['cogs']}",
            f"={c}{N['revenue']}*(1-INPUTS!{c}${I['gross_margin']})", fmt=NUM)
        put(ws, f"{c}{N['gross_profit']}",
            f"={c}{N['revenue']}-{c}{N['cogs']}", fmt=NUM)
        put(ws, f"{c}{N['sga']}",
            f"={c}{N['revenue']}*INPUTS!{c}${I['sga_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{N['ebitda']}",
            f"={c}{N['gross_profit']}-{c}{N['sga']}", fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{N['da']}",
            f"={c}{N['revenue']}*INPUTS!{c}${I['da_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{N['ebit']}", f"={c}{N['ebitda']}-{c}{N['da']}", fmt=NUM)
        put(ws, f"{c}{N['interest_expense']}",
            f"={open_debt}*INPUTS!$B$5+{open_rev}*INPUTS!$B$6", fmt=NUM)
        put(ws, f"{c}{N['interest_income']}",
            f"={open_cash}*INPUTS!$B$7", fmt=NUM)
        put(ws, f"{c}{N['pbt']}",
            f"={c}{N['ebit']}-{c}{N['interest_expense']}+{c}{N['interest_income']}",
            fmt=NUM)
        put(ws, f"{c}{N['tax']}",
            f"=MAX({c}{N['pbt']},0)*INPUTS!$B$4", fmt=NUM)
        put(ws, f"{c}{N['net_income']}",
            f"={c}{N['pbt']}-{c}{N['tax']}", fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{N['dividends']}",
            f"=MAX({c}{N['net_income']},0)*INPUTS!$B$9", fmt=NUM)
        put(ws, f"{c}{N['retained']}",
            f"={c}{N['net_income']}-{c}{N['dividends']}", fmt=NUM)

    year_widths(ws, nf, label_w=34)
    ws.freeze_panes = "B4"
    return ws


def build_balance(wb, spec):
    ws = wb.create_sheet("BALANCE")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    HL = col(1 + nh)
    I, H, B, N, C = (STMT["INPUTS"], STMT["HISTORICALS"], STMT["BALANCE"],
                     STMT["INCOME"], STMT["CASHFLOW"])

    ws["A1"] = "BALANCE SHEET — forecast"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Other assets and other liabilities are held flat from the last "
                "reported balance sheet; they are not forecast.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Line item"] + [str(y) for y in core.fcst_years(spec)])

    for key, lab in [("cash", "Cash and equivalents"),
                     ("ar", "Accounts receivable"), ("inventory", "Inventory"),
                     ("other_ca", "Other current assets"), ("ppe_net", "Net PP&E"),
                     ("other_assets", "Other assets (held flat)"),
                     ("total_assets", "Total assets"),
                     ("ap", "Accounts payable"), ("accrued", "Accrued liabilities"),
                     ("other_cl", "Other current liabilities (held flat)"),
                     ("debt", "Term debt"), ("revolver", "Revolver"),
                     ("other_lt_liab", "Other long-term liabilities (held flat)"),
                     ("equity", "Shareholders' equity"),
                     ("total_le", "Total liabilities and equity"),
                     ("balance", "Balance check (assets − L&E)"),
                     ("nwc", "Net working capital (memo)"),
                     ("net_debt", "Net debt (memo)")]:
        label(ws, B[key], lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        open_ppe = (f"HISTORICALS!${HL}${H['ppe_net']}" if t == 0
                    else f"{p}{B['ppe_net']}")
        open_debt = (f"HISTORICALS!${HL}${H['debt']}" if t == 0
                     else f"{p}{B['debt']}")
        open_rev = "INPUTS!$B$10" if t == 0 else f"{p}{B['revolver']}"
        open_eq = (f"HISTORICALS!${HL}${H['equity']}" if t == 0
                   else f"{p}{B['equity']}")

        put(ws, f"{c}{B['cash']}", f"=CASHFLOW!{c}{C['closing_cash']}", fmt=NUM)
        put(ws, f"{c}{B['ar']}",
            f"=INCOME!{c}{N['revenue']}/{D}*INPUTS!{c}${I['dso']}", fmt=NUM)
        put(ws, f"{c}{B['inventory']}",
            f"=INCOME!{c}{N['cogs']}/{D}*INPUTS!{c}${I['dio']}", fmt=NUM)
        put(ws, f"{c}{B['other_ca']}",
            f"=INCOME!{c}{N['revenue']}*INPUTS!{c}${I['other_ca_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{B['ppe_net']}",
            f"={open_ppe}-CASHFLOW!{c}{C['capex']}-INCOME!{c}{N['da']}", fmt=NUM)
        put(ws, f"{c}{B['other_assets']}",
            f"=HISTORICALS!${HL}${H['other_assets']}", fmt=NUM)
        put(ws, f"{c}{B['total_assets']}",
            f"=SUM({c}{B['cash']}:{c}{B['other_assets']})", fmt=NUM, fill=FILL_BAND)

        put(ws, f"{c}{B['ap']}",
            f"=INCOME!{c}{N['cogs']}/{D}*INPUTS!{c}${I['dpo']}", fmt=NUM)
        put(ws, f"{c}{B['accrued']}",
            f"=INCOME!{c}{N['revenue']}*INPUTS!{c}${I['accrued_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{B['other_cl']}",
            f"=HISTORICALS!${HL}${H['other_cl']}", fmt=NUM)
        put(ws, f"{c}{B['debt']}",
            f"={open_debt}+CASHFLOW!{c}{C['debt_repayment']}", fmt=NUM)
        put(ws, f"{c}{B['revolver']}",
            f"={open_rev}+CASHFLOW!{c}{C['revolver_draw']}"
            f"+CASHFLOW!{c}{C['revolver_repay']}", fmt=NUM)
        put(ws, f"{c}{B['other_lt_liab']}",
            f"=HISTORICALS!${HL}${H['other_lt_liab']}", fmt=NUM)
        put(ws, f"{c}{B['equity']}",
            f"={open_eq}+INCOME!{c}{N['retained']}", fmt=NUM)
        put(ws, f"{c}{B['total_le']}",
            f"=SUM({c}{B['ap']}:{c}{B['equity']})", fmt=NUM, fill=FILL_BAND)
        put(ws, f"{c}{B['balance']}",
            f"={c}{B['total_assets']}-{c}{B['total_le']}", fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{B['nwc']}",
            f"={c}{B['ar']}+{c}{B['inventory']}+{c}{B['other_ca']}"
            f"-{c}{B['ap']}-{c}{B['accrued']}", fmt=NUM)
        put(ws, f"{c}{B['net_debt']}",
            f"={c}{B['debt']}+{c}{B['revolver']}-{c}{B['cash']}", fmt=NUM)

    year_widths(ws, nf, label_w=38)
    ws.freeze_panes = "B4"
    return ws


def build_cashflow(wb, spec):
    ws = wb.create_sheet("CASHFLOW")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    HL = col(1 + nh)
    I, H, B, N, C = (STMT["INPUTS"], STMT["HISTORICALS"], STMT["BALANCE"],
                     STMT["INCOME"], STMT["CASHFLOW"])

    ws["A1"] = "CASH FLOW STATEMENT — forecast"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("The revolver is the plug: it draws to hold the minimum cash "
                "balance and sweeps surplus cash to repay itself.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Line item"] + [str(y) for y in core.fcst_years(spec)])

    for key, lab in [("net_income", "Net income"), ("da", "Add back: D&A"),
                     ("delta_wc", "Change in working capital"),
                     ("cfo", "Cash from operations"), ("capex", "Capex"),
                     ("cfi", "Cash from investing"),
                     ("debt_repayment", "Term debt repayment"),
                     ("dividends", "Dividends paid"),
                     ("opening_cash", "Opening cash"),
                     ("cash_before_revolver", "Cash before revolver"),
                     ("revolver_draw", "Revolver draw"),
                     ("revolver_repay", "Revolver repayment"),
                     ("cff", "Cash from financing"),
                     ("net_change", "Net change in cash"),
                     ("closing_cash", "Closing cash"),
                     ("unlevered_fcf", "Unlevered free cash flow (memo)"),
                     ("levered_fcf", "Levered free cash flow (memo)")]:
        label(ws, C[key], lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        prev_nwc = (f"HISTORICALS!${HL}${H['nwc']}" if t == 0
                    else f"BALANCE!{p}{B['nwc']}")
        open_debt = (f"HISTORICALS!${HL}${H['debt']}" if t == 0
                     else f"BALANCE!{p}{B['debt']}")
        open_rev = "INPUTS!$B$10" if t == 0 else f"BALANCE!{p}{B['revolver']}"
        open_cash = (f"HISTORICALS!${HL}${H['cash']}" if t == 0
                     else f"BALANCE!{p}{B['cash']}")

        put(ws, f"{c}{C['net_income']}", f"=INCOME!{c}{N['net_income']}", fmt=NUM)
        put(ws, f"{c}{C['da']}", f"=INCOME!{c}{N['da']}", fmt=NUM)
        put(ws, f"{c}{C['delta_wc']}",
            f"=-(BALANCE!{c}{B['nwc']}-{prev_nwc})", fmt=NUM)
        put(ws, f"{c}{C['cfo']}",
            f"={c}{C['net_income']}+{c}{C['da']}+{c}{C['delta_wc']}",
            fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{C['capex']}",
            f"=-INCOME!{c}{N['revenue']}*INPUTS!{c}${I['capex_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{C['cfi']}", f"={c}{C['capex']}", fmt=NUM)
        put(ws, f"{c}{C['debt_repayment']}",
            f"=-MIN(INPUTS!{c}${I['debt_repayment']},{open_debt})", fmt=NUM)
        put(ws, f"{c}{C['dividends']}", f"=-INCOME!{c}{N['dividends']}", fmt=NUM)
        put(ws, f"{c}{C['opening_cash']}", f"={open_cash}", fmt=NUM)
        put(ws, f"{c}{C['cash_before_revolver']}",
            f"={c}{C['opening_cash']}+{c}{C['cfo']}+{c}{C['cfi']}"
            f"+{c}{C['debt_repayment']}+{c}{C['dividends']}", fmt=NUM)
        put(ws, f"{c}{C['revolver_draw']}",
            f"=MAX(0,INPUTS!$B$8-{c}{C['cash_before_revolver']})", fmt=NUM)
        put(ws, f"{c}{C['revolver_repay']}",
            f"=-MIN({open_rev},MAX(0,{c}{C['cash_before_revolver']}-INPUTS!$B$8))",
            fmt=NUM)
        put(ws, f"{c}{C['cff']}",
            f"={c}{C['debt_repayment']}+{c}{C['dividends']}"
            f"+{c}{C['revolver_draw']}+{c}{C['revolver_repay']}", fmt=NUM)
        put(ws, f"{c}{C['net_change']}",
            f"={c}{C['cfo']}+{c}{C['cfi']}+{c}{C['cff']}", fmt=NUM)
        put(ws, f"{c}{C['closing_cash']}",
            f"={c}{C['opening_cash']}+{c}{C['net_change']}", fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{C['unlevered_fcf']}",
            f"=INCOME!{c}{N['ebit']}*(1-INPUTS!$B$4)+INCOME!{c}{N['da']}"
            f"+{c}{C['capex']}+{c}{C['delta_wc']}", fmt=NUM)
        put(ws, f"{c}{C['levered_fcf']}",
            f"={c}{C['cfo']}+{c}{C['cfi']}", fmt=NUM)

    year_widths(ws, nf, label_w=38)
    ws.freeze_panes = "B4"
    return ws


def build_ratios(wb, spec):
    """Historical and forecast ratios side by side -- the diagnostic read."""
    ws = wb.create_sheet("RATIOS")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    H, B, N, C, R = (STMT["HISTORICALS"], STMT["BALANCE"], STMT["INCOME"],
                     STMT["CASHFLOW"], STMT["RATIOS"])

    ws["A1"] = "RATIOS — history and forecast on one line, so the break shows"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("If a ratio steps sharply at the forecast boundary, the "
                "assumption behind it needs a reason.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Ratio"] + [str(y) for y in core.hist_years(spec)]
         + [str(y) for y in core.fcst_years(spec)])
    for k in range(nf):
        ws.cell(row=3, column=2 + nh + k).fill = FILL_BAND

    rows = [("revenue_growth", "Revenue growth", PCT),
            ("ebitda_growth", "EBITDA growth", PCT),
            ("gross_margin", "Gross margin", PCT),
            ("ebitda_margin", "EBITDA margin", PCT),
            ("ebit_margin", "EBIT margin", PCT),
            ("net_margin", "Net margin", PCT),
            ("capex_pct", "Capex % of revenue", PCT),
            ("dso", "Days sales outstanding", DAYS),
            ("dio", "Days inventory outstanding", DAYS),
            ("dpo", "Days payables outstanding", DAYS),
            ("ccc", "Cash conversion cycle (days)", DAYS),
            ("fixed_asset_turnover", "Fixed asset turnover", MULT),
            ("net_debt_ebitda", "Net debt / EBITDA", MULT),
            ("interest_cover", "EBIT / interest expense", MULT),
            ("roce", "Return on capital employed", PCT),
            ("roe", "Return on equity", PCT),
            ("retention", "Retention ratio", PCT),
            ("sustainable_growth", "Sustainable growth (retention x ROE)", PCT),
            ("cfo_sales", "CFO / revenue", PCT),
            ("cfo_assets", "CFO / total assets", PCT),
            ("cfo_debt", "CFO / total debt", PCT),
            ("fcf_conversion", "Levered FCF / EBITDA", PCT)]
    for key, lab, _ in rows:
        label(ws, R[key], lab)

    # historical columns
    for t in range(nh):
        c, p = col(2 + t), col(1 + t)
        def hist(row):
            return f"HISTORICALS!{c}{row}"
        put(ws, f"{c}{R['revenue_growth']}",
            "=NA()" if t == 0
            else f"=HISTORICALS!{c}{H['revenue']}/HISTORICALS!{p}{H['revenue']}-1",
            fmt=PCT)
        put(ws, f"{c}{R['gross_margin']}",
            f"={hist(H['gross_profit'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['ebitda_margin']}",
            f"={hist(H['ebitda'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['ebit_margin']}",
            f"={hist(H['ebit'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['ebitda_growth']}",
            "=NA()" if t == 0
            else f"=HISTORICALS!{c}{H['ebitda']}/HISTORICALS!{p}{H['ebitda']}-1",
            fmt=PCT)
        put(ws, f"{c}{R['net_margin']}",
            f"={hist(H['net_income'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['fixed_asset_turnover']}",
            f"={hist(H['revenue'])}/{hist(H['ppe_net'])}", fmt=MULT)
        put(ws, f"{c}{R['roce']}",
            f"={hist(H['ebit'])}/({hist(H['total_assets'])}"
            f"-{hist(H['ap'])}-{hist(H['accrued'])})", fmt=PCT)
        put(ws, f"{c}{R['retention']}", "=NA()", fmt=PCT)
        put(ws, f"{c}{R['sustainable_growth']}", "=NA()", fmt=PCT)
        put(ws, f"{c}{R['cfo_sales']}",
            f"={hist(H['cfo'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['cfo_assets']}",
            f"={hist(H['cfo'])}/{hist(H['total_assets'])}", fmt=PCT)
        put(ws, f"{c}{R['cfo_debt']}",
            f"={hist(H['cfo'])}/{hist(H['debt'])}", fmt=PCT)
        put(ws, f"{c}{R['capex_pct']}",
            f"={hist(H['capex'])}/{hist(H['revenue'])}", fmt=PCT)
        put(ws, f"{c}{R['dso']}",
            f"={hist(H['ar'])}/{hist(H['revenue'])}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['dio']}",
            f"={hist(H['inventory'])}/{hist(H['cogs'])}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['dpo']}",
            f"={hist(H['ap'])}/{hist(H['cogs'])}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['ccc']}",
            f"={c}{R['dso']}+{c}{R['dio']}-{c}{R['dpo']}", fmt=DAYS)
        put(ws, f"{c}{R['net_debt_ebitda']}",
            f"=({hist(H['debt'])}-{hist(H['cash'])})/{hist(H['ebitda'])}", fmt=MULT)
        put(ws, f"{c}{R['interest_cover']}",
            f"={hist(H['ebit'])}/{hist(H['interest_expense'])}", fmt=MULT)
        put(ws, f"{c}{R['roe']}",
            "=NA()" if t == 0 else
            f"={hist(H['net_income'])}/((HISTORICALS!{c}{H['equity']}"
            f"+HISTORICALS!{p}{H['equity']})/2)", fmt=PCT)
        put(ws, f"{c}{R['fcf_conversion']}", "=NA()", fmt=PCT)

    HL = col(1 + nh)
    for t in range(nf):
        c = col(2 + nh + t)
        fc, fp = col(2 + t), col(1 + t)
        prev_rev = (f"HISTORICALS!${HL}${H['revenue']}" if t == 0
                    else f"INCOME!{fp}{N['revenue']}")
        put(ws, f"{c}{R['revenue_growth']}",
            f"=INCOME!{fc}{N['revenue']}/{prev_rev}-1", fmt=PCT)
        put(ws, f"{c}{R['gross_margin']}",
            f"=INCOME!{fc}{N['gross_profit']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['ebitda_margin']}",
            f"=INCOME!{fc}{N['ebitda']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['ebit_margin']}",
            f"=INCOME!{fc}{N['ebit']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['net_margin']}",
            f"=INCOME!{fc}{N['net_income']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['capex_pct']}",
            f"=-CASHFLOW!{fc}{C['capex']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['dso']}",
            f"=BALANCE!{fc}{B['ar']}/INCOME!{fc}{N['revenue']}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['dio']}",
            f"=BALANCE!{fc}{B['inventory']}/INCOME!{fc}{N['cogs']}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['dpo']}",
            f"=BALANCE!{fc}{B['ap']}/INCOME!{fc}{N['cogs']}*{D}", fmt=DAYS)
        put(ws, f"{c}{R['ccc']}",
            f"={c}{R['dso']}+{c}{R['dio']}-{c}{R['dpo']}", fmt=DAYS)
        put(ws, f"{c}{R['net_debt_ebitda']}",
            f"=BALANCE!{fc}{B['net_debt']}/INCOME!{fc}{N['ebitda']}", fmt=MULT)
        put(ws, f"{c}{R['interest_cover']}",
            f"=INCOME!{fc}{N['ebit']}/INCOME!{fc}{N['interest_expense']}", fmt=MULT)
        put(ws, f"{c}{R['roe']}",
            f"=INCOME!{fc}{N['net_income']}/DUPONT!{c}{DUP['avg_equity']}", fmt=PCT)
        put(ws, f"{c}{R['fcf_conversion']}",
            f"=CASHFLOW!{fc}{C['levered_fcf']}/INCOME!{fc}{N['ebitda']}", fmt=PCT)
        prev_ebitda = (f"HISTORICALS!${HL}${H['ebitda']}" if t == 0
                       else f"INCOME!{fp}{N['ebitda']}")
        put(ws, f"{c}{R['ebitda_growth']}",
            f"=INCOME!{fc}{N['ebitda']}/{prev_ebitda}-1", fmt=PCT)
        put(ws, f"{c}{R['fixed_asset_turnover']}",
            f"=INCOME!{fc}{N['revenue']}/BALANCE!{fc}{B['ppe_net']}", fmt=MULT)
        put(ws, f"{c}{R['roce']}",
            f"=INCOME!{fc}{N['ebit']}/(BALANCE!{fc}{B['total_assets']}"
            f"-BALANCE!{fc}{B['ap']}-BALANCE!{fc}{B['accrued']})", fmt=PCT)
        put(ws, f"{c}{R['retention']}",
            f"=1-INCOME!{fc}{N['dividends']}/INCOME!{fc}{N['net_income']}", fmt=PCT)
        put(ws, f"{c}{R['sustainable_growth']}",
            f"={c}{R['retention']}*{c}{R['roe']}", fmt=PCT)
        put(ws, f"{c}{R['cfo_sales']}",
            f"=CASHFLOW!{fc}{C['cfo']}/INCOME!{fc}{N['revenue']}", fmt=PCT)
        put(ws, f"{c}{R['cfo_assets']}",
            f"=CASHFLOW!{fc}{C['cfo']}/BALANCE!{fc}{B['total_assets']}", fmt=PCT)
        put(ws, f"{c}{R['cfo_debt']}",
            f"=CASHFLOW!{fc}{C['cfo']}/(BALANCE!{fc}{B['debt']}"
            f"+BALANCE!{fc}{B['revolver']})", fmt=PCT)

    year_widths(ws, nh + nf, label_w=34, year_w=11)
    ws.freeze_panes = "B4"
    return ws



DUP = {"header": 3, "net_income": 4, "revenue": 5, "net_margin": 6,
       "avg_assets": 8, "asset_turnover": 9,
       "avg_equity": 11, "equity_multiplier": 12,
       "roe": 14, "roe_direct": 15, "roe_check": 16,
       "roa": 18}


def build_dupont(wb, spec):
    """ROE decomposed on average balances, history and forecast side by side."""
    ws = wb.create_sheet("DUPONT")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    H, B, N = STMT["HISTORICALS"], STMT["BALANCE"], STMT["INCOME"]
    HL = col(1 + nh)

    ws["A1"] = "DUPONT — where the return on equity actually comes from"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Margin, asset efficiency and leverage. Two companies with the "
                "same ROE can be entirely different businesses; this says which. "
                "Balances are averaged, so a year-end raise does not flatter the year.")
    ws["A2"].font = F_NOTE
    head(ws, DUP["header"], ["Component"] + [str(y) for y in core.hist_years(spec)]
         + [str(y) for y in core.fcst_years(spec)])
    for k in range(nf):
        ws.cell(row=DUP["header"], column=2 + nh + k).fill = FILL_BAND

    for key, lab in [("net_income", "Net income"), ("revenue", "Revenue"),
                     ("net_margin", "Net profit margin (A)"),
                     ("avg_assets", "Average total assets"),
                     ("asset_turnover", "Asset turnover (B)"),
                     ("avg_equity", "Average shareholders' equity"),
                     ("equity_multiplier", "Equity multiplier (C)"),
                     ("roe", "Return on equity (A x B x C)"),
                     ("roe_direct", "Return on equity (direct)"),
                     ("roe_check", "Difference"),
                     ("roa", "Return on assets (A x B)")]:
        label(ws, DUP[key], lab)

    # historical columns: the first has no opening balance, so averages start at year 2
    for t in range(nh):
        c, p = col(2 + t), col(1 + t)
        put(ws, f"{c}{DUP['net_income']}",
            f"=HISTORICALS!{c}{H['net_income']}", fmt=NUM)
        put(ws, f"{c}{DUP['revenue']}", f"=HISTORICALS!{c}{H['revenue']}", fmt=NUM)
        put(ws, f"{c}{DUP['net_margin']}",
            f"={c}{DUP['net_income']}/{c}{DUP['revenue']}", fmt=PCT)
        if t == 0:
            for key in ("avg_assets", "asset_turnover", "avg_equity",
                        "equity_multiplier", "roe", "roe_direct", "roe_check",
                        "roa"):
                put(ws, f"{c}{DUP[key]}", "=NA()", fmt=NUM)
            continue
        put(ws, f"{c}{DUP['avg_assets']}",
            f"=(HISTORICALS!{c}{H['total_assets']}"
            f"+HISTORICALS!{p}{H['total_assets']})/2", fmt=NUM)
        put(ws, f"{c}{DUP['asset_turnover']}",
            f"={c}{DUP['revenue']}/{c}{DUP['avg_assets']}", fmt=MULT)
        put(ws, f"{c}{DUP['avg_equity']}",
            f"=(HISTORICALS!{c}{H['equity']}+HISTORICALS!{p}{H['equity']})/2", fmt=NUM)
        put(ws, f"{c}{DUP['equity_multiplier']}",
            f"={c}{DUP['avg_assets']}/{c}{DUP['avg_equity']}", fmt=MULT)
        put(ws, f"{c}{DUP['roe']}",
            f"={c}{DUP['net_margin']}*{c}{DUP['asset_turnover']}"
            f"*{c}{DUP['equity_multiplier']}", fmt=PCT, fill=FILL_BAND)
        put(ws, f"{c}{DUP['roe_direct']}",
            f"={c}{DUP['net_income']}/{c}{DUP['avg_equity']}", fmt=PCT)
        put(ws, f"{c}{DUP['roe_check']}",
            f"={c}{DUP['roe']}-{c}{DUP['roe_direct']}", fmt="0.000000")
        put(ws, f"{c}{DUP['roa']}",
            f"={c}{DUP['net_margin']}*{c}{DUP['asset_turnover']}", fmt=PCT)

    for t in range(nf):
        c = col(2 + nh + t)
        fc, fp = col(2 + t), col(1 + t)
        prev_assets = (f"HISTORICALS!${HL}${H['total_assets']}" if t == 0
                       else f"BALANCE!{fp}{B['total_assets']}")
        prev_equity = (f"HISTORICALS!${HL}${H['equity']}" if t == 0
                       else f"BALANCE!{fp}{B['equity']}")
        put(ws, f"{c}{DUP['net_income']}", f"=INCOME!{fc}{N['net_income']}", fmt=NUM)
        put(ws, f"{c}{DUP['revenue']}", f"=INCOME!{fc}{N['revenue']}", fmt=NUM)
        put(ws, f"{c}{DUP['net_margin']}",
            f"={c}{DUP['net_income']}/{c}{DUP['revenue']}", fmt=PCT)
        put(ws, f"{c}{DUP['avg_assets']}",
            f"=(BALANCE!{fc}{B['total_assets']}+{prev_assets})/2", fmt=NUM)
        put(ws, f"{c}{DUP['asset_turnover']}",
            f"={c}{DUP['revenue']}/{c}{DUP['avg_assets']}", fmt=MULT)
        put(ws, f"{c}{DUP['avg_equity']}",
            f"=(BALANCE!{fc}{B['equity']}+{prev_equity})/2", fmt=NUM)
        put(ws, f"{c}{DUP['equity_multiplier']}",
            f"={c}{DUP['avg_assets']}/{c}{DUP['avg_equity']}", fmt=MULT)
        put(ws, f"{c}{DUP['roe']}",
            f"={c}{DUP['net_margin']}*{c}{DUP['asset_turnover']}"
            f"*{c}{DUP['equity_multiplier']}", fmt=PCT, fill=FILL_OUT)
        put(ws, f"{c}{DUP['roe_direct']}",
            f"={c}{DUP['net_income']}/{c}{DUP['avg_equity']}", fmt=PCT)
        put(ws, f"{c}{DUP['roe_check']}",
            f"={c}{DUP['roe']}-{c}{DUP['roe_direct']}", fmt="0.000000")
        put(ws, f"{c}{DUP['roa']}",
            f"={c}{DUP['net_margin']}*{c}{DUP['asset_turnover']}", fmt=PCT)

    note(ws, DUP["roe_check"], 2 + nh + nf,
         "Must be zero: the decomposition is an identity, so if it drifts the "
         "components are not the ones being multiplied")
    year_widths(ws, nh + nf, label_w=34, year_w=11)
    ws.freeze_panes = "B4"
    return ws


def build_checks(wb, spec):
    """In-workbook assertions.

    Every range is sheet-qualified: an unqualified range silently points at
    this sheet's own empty cells and the check passes against nothing.
    MAX(ABS(range)) needs array evaluation, so each one is wrapped in
    SUMPRODUCT, which forces it in both Excel and LibreOffice.
    """
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    HL, FL = col(1 + nh), col(1 + nf)
    H, B, N, C, I = (STMT["HISTORICALS"], STMT["BALANCE"], STMT["INCOME"],
                     STMT["CASHFLOW"], STMT["INPUTS"])

    def rng(sheet, row):
        return f"{sheet}!B{row}:{FL}{row}"

    def hrng(row):
        return f"HISTORICALS!B{row}:{HL}{row}"

    return checks_sheet(wb, [
        ("Forecast balance sheet balances every year",
         f'=IF(SUMPRODUCT(MAX(ABS({rng("BALANCE", B["balance"])})))<0.01,'
         f'"PASS","FAIL")',
         "Assets = liabilities + equity in every projected year. The single "
         "most important check in a three-statement model"),
        ("Reported balance sheet balanced every year",
         f'=IF(SUMPRODUCT(MAX(ABS({hrng(H["balance"])})))<0.01,"PASS","FAIL")',
         "The history was transcribed correctly"),
        ("Cash never falls below the minimum",
         f'=IF(MIN({rng("BALANCE", B["cash"])})>=INPUTS!$B$8-0.01,"PASS","FAIL")',
         "The revolver is doing its job as the plug"),
        ("Closing cash ties to the cash flow statement",
         f'=IF(SUMPRODUCT(MAX(ABS({rng("CASHFLOW", C["closing_cash"])}'
         f'-{rng("BALANCE", B["cash"])})))<0.01,"PASS","FAIL")',
         "Balance sheet cash is the cash flow statement's closing line"),
        ("Equity rolls forward correctly",
         f'=IF(ABS(BALANCE!B{B["equity"]}-(HISTORICALS!${HL}${H["equity"]}'
         f'+INCOME!B{N["retained"]}))<0.01,"PASS","FAIL")',
         "Opening equity + net income - dividends = closing equity"),
        ("PP&E rolls forward correctly",
         f'=IF(ABS(BALANCE!B{B["ppe_net"]}-(HISTORICALS!${HL}${H["ppe_net"]}'
         f'-CASHFLOW!B{C["capex"]}-INCOME!B{N["da"]}))<0.01,"PASS","FAIL")',
         "Opening PP&E + capex - depreciation = closing PP&E"),
        ("Cash from operations identity",
         f'=IF(SUMPRODUCT(MAX(ABS({rng("CASHFLOW", C["cfo"])}'
         f'-{rng("CASHFLOW", C["net_income"])}'
         f'-{rng("CASHFLOW", C["da"])}'
         f'-{rng("CASHFLOW", C["delta_wc"])})))<0.01,"PASS","FAIL")',
         "CFO = net income + D&A + change in working capital"),
        ("Term debt never goes negative",
         f'=IF(MIN({rng("BALANCE", B["debt"])})>=-0.01,"PASS","FAIL")',
         "Scheduled repayment never exceeds the balance outstanding"),
        ("Interest charged on opening balances",
         f'=IF(ABS(INCOME!B{N["interest_expense"]}'
         f'-(HISTORICALS!${HL}${H["debt"]}*INPUTS!$B$5'
         f'+INPUTS!$B$10*INPUTS!$B$6))<0.01,"PASS","FAIL")',
         "No circular reference: year 1 interest uses the last reported debt"),
        ("Revenue growth matches the input",
         f'=IF(ABS(INCOME!C{N["revenue"]}/INCOME!B{N["revenue"]}-1'
         f'-INPUTS!C{I["rev_growth"]})<0.0001,"PASS","FAIL")',
         "The statements are actually driven by the assumptions"),
        ("Working capital days match the inputs",
         f'=IF(ABS(BALANCE!B{B["ar"]}/INCOME!B{N["revenue"]}*{D}'
         f'-INPUTS!B{I["dso"]})<0.01,"PASS","FAIL")',
         "Receivables are genuinely driven by the DSO assumption"),
        ("DuPont decomposition is an identity",
         f'=IF(SUMPRODUCT(MAX(ABS(DUPONT!{col(2 + nh)}{DUP["roe_check"]}:'
         f'{col(1 + nh + nf)}{DUP["roe_check"]})))<0.000001,"PASS","FAIL")',
         "Margin x asset turnover x equity multiplier equals ROE computed "
         "directly, in every forecast year"),
        ("Year-1 EBITDA margin at or above the historical trough",
         f'=IF(INCOME!B{N["ebitda"]}/INCOME!B{N["revenue"]}'
         f'>=HISTORICALS!$B${H["em_min"]},"PASS","REVIEW")',
         "The forecast is not starting below anything the company has lived "
         "through; REVIEW means justify it, not hide it"),
    ])

def build(spec, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    readme_sheet(wb, spec, "Three-statement operating model", [
        ("History", f"{core.n_hist(spec)} reported years of income statement, "
                    "balance sheet and operating cash flow"),
        ("DUPONT tab", "Return on equity split into margin, asset efficiency and "
                       "leverage, on average balances"),
        ("Forecast", f"{core.n_fcst(spec)} years, fully linked"),
        ("Interest", "Charged on opening balances, so the workbook needs no "
                     "iterative calculation. It slightly understates interest "
                     "in a year of heavy drawdown."),
        ("Cash plug", "A revolver draws to hold the minimum cash balance and "
                      "sweeps surplus cash to repay itself."),
        ("Held flat", "Other assets, other current liabilities and other "
                      "long-term liabilities are carried at their last reported "
                      "value rather than forecast."),
    ])
    build_inputs(wb, spec)
    build_historicals(wb, spec)
    build_income(wb, spec)
    build_balance(wb, spec)
    build_cashflow(wb, spec)
    build_ratios(wb, spec)
    build_dupont(wb, spec)
    build_checks(wb, spec)
    wb.save(out_path)
    return out_path
