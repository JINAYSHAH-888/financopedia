#!/usr/bin/env python3
"""The derivation tabs: WACC build-up, intrinsic growth, relative pricing,
scenarios and the football field.

These are the tabs that turn asserted numbers into arguments. A bare
"beta: 1.1" cannot be checked; a peer set unlevered, re-levered at a target
structure and blended with a regression can be.
"""

from __future__ import annotations

from openpyxl.chart import BarChart, Reference

import dcf_core as core
import valuation_core as vc
from xl_style import (CUR, F_INPUT, F_NOTE, F_SUB, F_TITLE, FILL_BAND, year_widths,
                      FILL_OUT, MULT, NUM, PCT, PCT2, col, head, label, note,
                      put, widths)

# --------------------------------------------------------------------------
# cell maps
# --------------------------------------------------------------------------

W = {
    "peer_header": 4, "peer_first": 5,          # + n peers
    # the rest are offsets resolved by wacc_rows()
}


def wacc_rows(np_):
    r = {"peer_header": 4, "peer_first": 5}
    r["peer_last"] = 4 + np_
    r["median"] = r["peer_last"] + 1
    b = r["median"] + 2
    r["reg_header"] = b
    r["raw_beta"] = b + 1
    r["blume_weight"] = b + 2
    r["market_beta"] = b + 3
    r["adjusted_beta"] = b + 4
    r["r_squared"] = b + 5
    c = b + 7
    r["sel_header"] = c
    r["target_de"] = c + 1
    r["relevered"] = c + 2
    r["beta_source"] = c + 3
    r["beta_used"] = c + 4
    d = c + 6
    r["erp_header"] = d
    r["mr_header"] = d + 1
    r["mr_first"] = d + 2
    return r


def erp_rows(r, n_years):
    r = dict(r)
    r["mr_last"] = r["mr_first"] + n_years - 1
    e = r["mr_last"] + 2
    r["mr_avg"] = e
    r["div_yield"] = e + 1
    r["market_total"] = e + 2
    r["risk_free"] = e + 3
    r["erp"] = e + 4
    f = e + 6
    r["kd_header"] = f
    r["interest_expense"] = f + 1
    r["total_debt"] = f + 2
    r["pre_tax_kd"] = f + 3
    r["tax_rate"] = f + 4
    r["after_tax_kd"] = f + 5
    g = f + 7
    r["wacc_header"] = g
    r["cost_of_equity"] = g + 1
    r["equity_weight"] = g + 2
    r["debt_weight"] = g + 3
    r["wacc"] = g + 4
    return r


IG = {"header": 3, "ppe": 4, "nwc": 5, "invested_capital": 6, "ebit": 7,
      "nopat": 8, "roic": 9,
      "rr_header": 11, "net_capex": 12, "delta_nwc": 13, "reinvestment": 14,
      "nopat_latest": 15, "reinvestment_rate": 16,
      "g_header": 18, "roic_used": 19, "intrinsic_growth": 20,
      "assumed_growth": 21, "terminal_growth": 22, "verdict": 23}


def rel_rows(np_):
    r = {"header": 5, "first": 6}
    r["last"] = 5 + np_
    s = r["last"] + 2
    r["high"] = s
    r["p75"] = s + 1
    r["mean"] = s + 2
    r["median"] = s + 3
    r["p25"] = s + 4
    r["low"] = s + 5
    t = r["low"] + 2
    r["subject_header"] = t
    r["sub_revenue"] = t + 1
    r["sub_ebitda"] = t + 2
    r["sub_net_income"] = t + 3
    r["sub_net_debt"] = t + 4
    r["sub_shares"] = t + 5
    u = r["sub_shares"] + 2
    r["implied_header"] = u
    r["implied_p25"] = u + 1
    r["implied_median"] = u + 2
    r["implied_p75"] = u + 3
    return r


FF = {"header": 4, "first": 5, "price": 12}


def scen_rows(nf):
    r = {"header": 4, "bear": 5, "base": 6, "bull": 7, "helper": 10}
    r["block"] = {}
    for k, name in enumerate(("bear", "base", "bull")):
        r["block"][name] = r["helper"] + 2 + k * 12
    return r


# --------------------------------------------------------------------------
# WACC
# --------------------------------------------------------------------------


def build_wacc(wb, spec, raw_beta_ref=None):
    w = spec["wacc_build"]
    peers = w.get("peers", [])
    mr = w.get("market_returns", {})
    years = mr.get("years", [])
    R = erp_rows(wacc_rows(len(peers)), len(years))
    ws = wb.create_sheet("WACC")

    ws["A1"] = "WACC — derived, not asserted"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Beta from a peer set and a regression; the equity risk premium "
                "from a market return history; the cost of debt from interest "
                "actually paid. Blue cells are the sourced inputs.")
    ws["A2"].font = F_NOTE

    head(ws, R["peer_header"],
         ["Peer", "Total debt", "Total equity", "Tax rate", "D/E", "D/(D+E)",
          "Levered beta", "Unlevered beta", "Source"])
    for k, p in enumerate(peers):
        r = R["peer_first"] + k
        label(ws, r, p["name"])
        put(ws, f"B{r}", p["total_debt"], font=F_INPUT, fmt=NUM)
        put(ws, f"C{r}", p["total_equity"], font=F_INPUT, fmt=NUM)
        put(ws, f"D{r}", p["tax_rate"], font=F_INPUT, fmt=PCT)
        put(ws, f"E{r}", f"=B{r}/C{r}", fmt="0.00")
        put(ws, f"F{r}", f"=B{r}/(B{r}+C{r})", fmt=PCT)
        put(ws, f"G{r}", p["levered_beta"], font=F_INPUT, fmt="0.000")
        put(ws, f"H{r}", f"=G{r}/(1+(1-D{r})*E{r})", fmt="0.000")
        note(ws, r, 9, p.get("source", ""))
    label(ws, R["median"], "Median")
    for c in ("E", "F", "G", "H"):
        put(ws, f"{c}{R['median']}",
            f"=MEDIAN({c}{R['peer_first']}:{c}{R['peer_last']})",
            fmt="0.000" if c in ("G", "H") else (PCT if c == "F" else "0.00"),
            fill=FILL_BAND)
    note(ws, R["median"], 9,
         "Unlevering strips each peer's financing out of its observed beta, "
         "leaving business risk only (Hamada)")

    ws.cell(row=R["reg_header"], column=1,
            value="Regression beta, Blume-adjusted").font = F_TITLE
    reg = w.get("regression_beta") or {}
    for r, lab, val, fmt, is_input in [
        (R["raw_beta"], "Raw regression beta", reg.get("raw_beta", 1.0), "0.000",
         raw_beta_ref is None),
        (R["blume_weight"], "Weight on the raw beta",
         reg.get("blume_weight_raw", 0.75), PCT, True),
        (R["market_beta"], "Market beta (shrink target)", 1.0, "0.000", True),
        (R["adjusted_beta"], "Adjusted beta", None, "0.000", False),
        (R["r_squared"], "R² of the regression", reg.get("r_squared", 0.0),
         "0.00", True),
    ]:
        label(ws, r, lab)
        if r == R["raw_beta"] and raw_beta_ref is not None:
            put(ws, f"B{r}", f"={raw_beta_ref}", fmt=fmt)
        elif r == R["r_squared"] and raw_beta_ref is not None:
            put(ws, f"B{r}", f"=BETA_REGRESSION!H{br_stats_rows()['r_squared']}",
                fmt=fmt)
        elif is_input:
            put(ws, f"B{r}", val, font=F_INPUT, fmt=fmt)
        else:
            put(ws, f"B{r}",
                f"=B{R['raw_beta']}*B{R['blume_weight']}"
                f"+B{R['market_beta']}*(1-B{R['blume_weight']})",
                fmt=fmt, fill=FILL_BAND)
    note(ws, R["raw_beta"], 3, reg.get("source", ""))
    note(ws, R["adjusted_beta"], 3,
         "Blume: betas drift toward 1 over time, so the raw estimate is shrunk "
         "toward the market")
    note(ws, R["r_squared"], 3,
         "How much of the share's movement the index explains. A low figure is "
         "a reason to lean on the peer-derived beta instead")

    ws.cell(row=R["sel_header"], column=1,
            value="Beta used in the cost of equity").font = F_TITLE
    label(ws, R["target_de"], "Target debt / equity")
    put(ws, f"B{R['target_de']}",
        w.get("target_debt_equity", 0.0), font=F_INPUT, fmt="0.00")
    label(ws, R["relevered"], "Peer beta re-levered at the target structure")
    put(ws, f"B{R['relevered']}",
        f"=H{R['median']}*(1+(1-B{R['tax_rate']})*B{R['target_de']})", fmt="0.000")
    label(ws, R["beta_source"], "Beta source (1 comps, 2 regression, 3 average)")
    put(ws, f"B{R['beta_source']}",
        {"comps": 1, "regression": 2, "average": 3}[w["beta_source"]],
        font=F_INPUT, fmt="0")
    label(ws, R["beta_used"], "Beta used")
    put(ws, f"B{R['beta_used']}",
        f"=CHOOSE(B{R['beta_source']},B{R['relevered']},B{R['adjusted_beta']},"
        f"AVERAGE(B{R['relevered']},B{R['adjusted_beta']}))",
        fmt="0.000", fill=FILL_OUT)

    ws.cell(row=R["erp_header"], column=1,
            value="Equity risk premium from the market record").font = F_TITLE
    head(ws, R["mr_header"], ["Year", "Market return"])
    for k, y in enumerate(years):
        r = R["mr_first"] + k
        put(ws, f"A{r}", y, font=F_INPUT, fmt="0")
        put(ws, f"B{r}", mr["returns"][k], font=F_INPUT, fmt=PCT)
    for r, lab, formula, fmt, val in [
        (R["mr_avg"], f"Average annual price return ({len(years)} years)",
         f"=AVERAGE(B{R['mr_first']}:B{R['mr_last']})", PCT, None),
        (R["div_yield"], "Dividend yield", None, PCT, mr.get("dividend_yield", 0.0)),
        (R["market_total"], "Total market return",
         f"=B{R['mr_avg']}+B{R['div_yield']}", PCT, None),
        (R["risk_free"], "Risk-free rate", None, PCT, w["risk_free"]),
        (R["erp"], "Equity risk premium",
         f"=B{R['market_total']}-B{R['risk_free']}", PCT, None),
    ]:
        label(ws, r, lab)
        if formula:
            put(ws, f"B{r}", formula, fmt=fmt,
                fill=FILL_OUT if r == R["erp"] else None)
        else:
            put(ws, f"B{r}", val, font=F_INPUT, fmt=fmt)
    note(ws, R["mr_avg"], 3, mr.get("source", ""))
    note(ws, R["erp"], 3,
         "A premium derived from a long record rather than a round number; the "
         "window length is itself an assumption")

    ws.cell(row=R["kd_header"], column=1, value="Cost of debt").font = F_TITLE
    cod = w["cost_of_debt"]
    uses_ratio = cod.get("method") == "interest_over_debt"
    for r, lab, val, formula, fmt in [
        (R["interest_expense"], "Interest expense",
         cod.get("interest_expense", 0.0), None, NUM),
        (R["total_debt"], "Gross debt", cod.get("total_debt", 0.0), None, NUM),
        (R["pre_tax_kd"], "Pre-tax cost of debt", cod.get("rate", 0.0),
         f"=B{R['interest_expense']}/B{R['total_debt']}" if uses_ratio else None, PCT),
        (R["tax_rate"], "Marginal tax rate", w["marginal_tax_rate"], None, PCT),
        (R["after_tax_kd"], "After-tax cost of debt", None,
         f"=B{R['pre_tax_kd']}*(1-B{R['tax_rate']})", PCT),
    ]:
        label(ws, r, lab)
        if formula:
            put(ws, f"B{r}", formula, fmt=fmt)
        else:
            put(ws, f"B{r}", val, font=F_INPUT, fmt=fmt)
    note(ws, R["pre_tax_kd"], 3, cod.get("source", ""))

    ws.cell(row=R["wacc_header"], column=1,
            value="Weighted average cost of capital").font = F_TITLE
    for r, lab, formula, fmt in [
        (R["cost_of_equity"], "Cost of equity",
         f"=B{R['risk_free']}+B{R['beta_used']}*B{R['erp']}", PCT2),
        (R["equity_weight"], "Equity weight",
         f"=1-B{R['debt_weight']}", PCT),
        (R["debt_weight"], "Debt weight",
         f"=B{R['target_de']}/(1+B{R['target_de']})", PCT),
        (R["wacc"], "WACC",
         f"=B{R['equity_weight']}*B{R['cost_of_equity']}"
         f"+B{R['debt_weight']}*B{R['after_tax_kd']}", PCT2),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt,
            fill=FILL_OUT if r == R["wacc"] else None)

    widths(ws, {"A": 42, "B": 15, "C": 15, "D": 12, "E": 10, "F": 11, "G": 13,
                "H": 14, "I": 54})
    return ws, R


# --------------------------------------------------------------------------
# intrinsic growth
# --------------------------------------------------------------------------


def build_intrinsic(wb, spec, hist_rows, wacc_tax_ref):
    ws = wb.create_sheet("INTRINSIC_GROWTH")
    nh = core.n_hist(spec)
    H = hist_rows
    last, prev = col(1 + nh), col(nh)

    ws["A1"] = "INTRINSIC GROWTH — what the business can fund"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("g = reinvestment rate × ROIC. Not a forecast — a benchmark to "
                "judge the forecast against.")
    ws["A2"].font = F_NOTE
    head(ws, IG["header"], ["Metric"] + [str(y) for y in core.hist_years(spec)])

    for key, lab in [("ppe", "Net PP&E"), ("nwc", "Net working capital"),
                     ("invested_capital", "Invested capital"),
                     ("ebit", "EBIT"), ("nopat", "NOPAT"), ("roic", "ROIC")]:
        label(ws, IG[key], lab)
    for t in range(nh):
        c = col(2 + t)
        put(ws, f"{c}{IG['ppe']}", f"=HISTORICALS!{c}{H['ppe_net']}", fmt=NUM)
        put(ws, f"{c}{IG['nwc']}", f"=HISTORICALS!{c}{H['nwc_balance']}", fmt=NUM)
        put(ws, f"{c}{IG['invested_capital']}",
            f"={c}{IG['ppe']}+{c}{IG['nwc']}", fmt=NUM)
        put(ws, f"{c}{IG['ebit']}", f"=HISTORICALS!{c}{H['ebit']}", fmt=NUM)
        put(ws, f"{c}{IG['nopat']}",
            f"={c}{IG['ebit']}*(1-{wacc_tax_ref})", fmt=NUM)
        put(ws, f"{c}{IG['roic']}",
            f"={c}{IG['nopat']}/{c}{IG['invested_capital']}", fmt=PCT,
            fill=FILL_BAND)

    ws.cell(row=IG["rr_header"], column=1,
            value="Reinvestment rate (latest year)").font = F_TITLE
    for r, lab, formula, fmt in [
        (IG["net_capex"], "Net capex (capex less D&A)",
         f"=HISTORICALS!{last}{H['capex']}-HISTORICALS!{last}{H['da']}", NUM),
        (IG["delta_nwc"], "Change in working capital",
         f"={last}{IG['nwc']}-{prev}{IG['nwc']}", NUM),
        (IG["reinvestment"], "Reinvestment",
         f"=B{IG['net_capex']}+B{IG['delta_nwc']}", NUM),
        (IG["nopat_latest"], "NOPAT", f"={last}{IG['nopat']}", NUM),
        (IG["reinvestment_rate"], "Reinvestment rate",
         f"=B{IG['reinvestment']}/B{IG['nopat_latest']}", PCT),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt,
            fill=FILL_BAND if r == IG["reinvestment_rate"] else None)

    ws.cell(row=IG["g_header"], column=1,
            value="Growth the company can fund").font = F_TITLE
    for r, lab, formula, fmt, fill in [
        (IG["roic_used"], "ROIC (latest)", f"={last}{IG['roic']}", PCT, None),
        (IG["intrinsic_growth"], "Intrinsic growth (RR × ROIC)",
         f"=B{IG['reinvestment_rate']}*B{IG['roic_used']}", PCT, FILL_OUT),
        (IG["assumed_growth"], "Year-1 revenue growth assumed on INPUTS",
         "=INPUTS!B20", PCT, None),
        (IG["terminal_growth"], "Terminal growth assumed on INPUTS",
         "=INPUTS!$B$11", PCT, None),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt, fill=fill)

    label(ws, IG["verdict"], "Assumed growth vs what the company can fund")
    put(ws, f"B{IG['verdict']}",
        f'=IF(B{IG["assumed_growth"]}<=B{IG["intrinsic_growth"]},'
        f'"Within intrinsic growth",'
        f'"Above intrinsic growth — needs margin expansion, better capital '
        f'efficiency, or a stated reason")', fill=FILL_BAND)
    note(ws, IG["verdict"], 3,
         "Growing faster than RR × ROIC means funding growth from somewhere "
         "the record does not show. That is allowed; it is not free")
    note(ws, IG["terminal_growth"], 3,
         "Terminal growth above intrinsic growth assumes the company improves "
         "forever — the strongest claim in any DCF")

    ws.column_dimensions["A"].width = 44
    for k in range(nh):
        ws.column_dimensions[col(2 + k)].width = 12
    ws.column_dimensions[col(2 + nh)].width = 60
    return ws


# --------------------------------------------------------------------------
# relative pricing
# --------------------------------------------------------------------------


def quartile_formula(rng: str, n: int, q: int) -> str:
    """QUARTILE.EXC written out longhand.

    The newer Excel statistical functions do not round-trip reliably through
    openpyxl and LibreOffice, and a comp sheet that silently returns the wrong
    percentile is worse than no percentile. SMALL and INT are universal.
    """
    h = q * (n + 1) / 4.0
    lo = int(h)
    frac = h - lo
    if lo < 1:
        return f"=MIN({rng})"
    if lo >= n:
        return f"=MAX({rng})"
    return (f"=SMALL({rng},{lo})+{frac}*(SMALL({rng},{lo + 1})"
            f"-SMALL({rng},{lo}))")


def build_relative(wb, spec):
    peers = spec["relative_peers"]
    sub = spec["subject"]
    np_ = len(peers)
    R = rel_rows(np_)
    ws = wb.create_sheet("RELATIVE_PRICING")

    ws["A1"] = "RELATIVE PRICING — the whole peer distribution"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("A median alone hides the spread. The quartiles are what the "
                "football field is built from.")
    ws["A2"].font = F_NOTE

    head(ws, R["header"],
         ["Peer", "Share price", "Shares", "Equity value", "Total debt", "Cash",
          "Net debt", "Enterprise value", "Revenue", "EBITDA", "Net income",
          "EV/Revenue", "EV/EBITDA", "P/E", "Source"])
    for k, p in enumerate(peers):
        r = R["first"] + k
        label(ws, r, p["name"])
        put(ws, f"B{r}", p["share_price"], font=F_INPUT, fmt=CUR)
        put(ws, f"C{r}", p["shares_out"], font=F_INPUT, fmt=NUM)
        put(ws, f"D{r}", f"=B{r}*C{r}", fmt=NUM)
        put(ws, f"E{r}", p["total_debt"], font=F_INPUT, fmt=NUM)
        put(ws, f"F{r}", p["cash"], font=F_INPUT, fmt=NUM)
        put(ws, f"G{r}", f"=E{r}-F{r}", fmt=NUM)
        put(ws, f"H{r}", f"=D{r}+G{r}", fmt=NUM)
        put(ws, f"I{r}", p["revenue"], font=F_INPUT, fmt=NUM)
        put(ws, f"J{r}", p["ebitda"], font=F_INPUT, fmt=NUM)
        put(ws, f"K{r}", p["net_income"], font=F_INPUT, fmt=NUM)
        put(ws, f"L{r}", f"=H{r}/I{r}", fmt=MULT)
        put(ws, f"M{r}", f"=H{r}/J{r}", fmt=MULT)
        put(ws, f"N{r}", f"=D{r}/K{r}", fmt=MULT)
        note(ws, r, 15, p.get("source", ""))

    for key, lab in [("high", "High"), ("p75", "75th percentile"),
                     ("mean", "Average"), ("median", "Median"),
                     ("p25", "25th percentile"), ("low", "Low")]:
        label(ws, R[key], lab)
        for c in ("L", "M", "N"):
            rng = f"{c}{R['first']}:{c}{R['last']}"
            formula = {"high": f"=MAX({rng})", "low": f"=MIN({rng})",
                       "mean": f"=AVERAGE({rng})",
                       "median": f"=MEDIAN({rng})",
                       "p75": quartile_formula(rng, np_, 3),
                       "p25": quartile_formula(rng, np_, 1)}[key]
            put(ws, f"{c}{R[key]}", formula, fmt=MULT,
                fill=FILL_BAND if key in ("p25", "median", "p75") else None)

    ws.cell(row=R["subject_header"], column=1,
            value="The subject company").font = F_TITLE
    for r, lab, key, fmt in [
        (R["sub_revenue"], "Revenue (NTM)", "revenue", NUM),
        (R["sub_ebitda"], "EBITDA (NTM)", "ebitda", NUM),
        (R["sub_net_income"], "Net income (NTM)", "net_income", NUM),
        (R["sub_net_debt"], "Net debt", "net_debt", NUM),
        (R["sub_shares"], "Shares outstanding", "shares_out", NUM),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", sub[key], font=F_INPUT, fmt=fmt)
    note(ws, R["sub_revenue"], 3,
         "On the same basis as the peers above — forward figures against "
         "forward multiples, or the comparison is meaningless")

    ws.cell(row=R["implied_header"], column=1,
            value="Implied value per share").font = F_TITLE
    head(ws, R["implied_header"], ["", "", "", "", "", "", "", "", "", "", "",
                                   "EV/Revenue", "EV/EBITDA", "P/E"])
    ws.cell(row=R["implied_header"], column=1,
            value="Implied value per share").font = F_TITLE
    for key, lab in [("implied_p25", "At the 25th percentile"),
                     ("implied_median", "At the median"),
                     ("implied_p75", "At the 75th percentile")]:
        stat = {"implied_p25": "p25", "implied_median": "median",
                "implied_p75": "p75"}[key]
        label(ws, R[key], lab)
        put(ws, f"L{R[key]}",
            f"=(L{R[stat]}*B{R['sub_revenue']}-B{R['sub_net_debt']})"
            f"/B{R['sub_shares']}", fmt=CUR)
        put(ws, f"M{R[key]}",
            f"=(M{R[stat]}*B{R['sub_ebitda']}-B{R['sub_net_debt']})"
            f"/B{R['sub_shares']}", fmt=CUR,
            fill=FILL_OUT if stat == "median" else None)
        put(ws, f"N{R[key]}",
            f"=N{R[stat]}*B{R['sub_net_income']}/B{R['sub_shares']}", fmt=CUR)
    note(ws, R["implied_p75"], 15,
         "P/E needs no net debt adjustment — earnings are already after interest")

    widths(ws, {"A": 30, "B": 13, "C": 11, "D": 14, "E": 12, "F": 11, "G": 12,
                "H": 16, "I": 12, "J": 12, "K": 12, "L": 13, "M": 13, "N": 13,
                "O": 46})
    return ws, R


# --------------------------------------------------------------------------
# scenarios
# --------------------------------------------------------------------------


def build_scenarios(wb, spec, cells):
    """Bear / base / bull, each a full formula-driven projection."""
    ws = wb.create_sheet("SCENARIOS")
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    H, I, V = cells["HISTORICALS"], cells["INPUTS"], cells["VALUATION"]
    hist_last, flast = col(1 + nh), col(1 + nf)
    S = scen_rows(nf)

    ws["A1"] = "SCENARIOS — bear, base and bull"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("A parallel shift to the growth path, set by the scenario delta "
                "on INPUTS. Each case is a full projection, not a haircut to "
                "the answer.")
    ws["A2"].font = F_NOTE
    head(ws, S["header"], ["Case", "Growth shift", "Value per share"])

    for name, sign in (("bear", -1), ("base", 0), ("bull", 1)):
        r = S[name]
        base = S["block"][name]
        label(ws, r, name.capitalize())
        put(ws, f"B{r}", f"={sign}*INPUTS!$B$19", fmt=PCT)
        put(ws, f"C{r}", f"=B{base + 11}", fmt=CUR,
            fill=FILL_OUT if name == "base" else None)

    ws.cell(row=S["helper"], column=1,
            value="Helper projections — one per case, driven off INPUTS with "
                  "that case's shift applied").font = F_NOTE

    for name, sign in (("bear", -1), ("base", 0), ("bull", 1)):
        base = S["block"][name]
        ws.cell(row=base, column=1,
                value=f"{name.capitalize()} case").font = F_SUB
        shift = f"+({sign}*INPUTS!$B$19)" if sign else ""
        rows = {"rev": base + 1, "ebit": base + 2, "nopat": base + 3,
                "da": base + 4, "capex": base + 5, "nwc": base + 6,
                "dnwc": base + 7, "fcf": base + 8, "pv": base + 9,
                "ev": base + 10, "vps": base + 11}
        for key, lab in [("rev", "Revenue"), ("ebit", "EBIT"), ("nopat", "NOPAT"),
                         ("da", "D&A"), ("capex", "Capex"), ("nwc", "NWC"),
                         ("dnwc", "Change in NWC"), ("fcf", "Unlevered FCF"),
                         ("pv", "PV of FCF")]:
            ws.cell(row=rows[key], column=1, value="    " + lab).font = F_NOTE

        for t in range(nf):
            c, p = col(2 + t), col(1 + t)
            prev_rev = (f"HISTORICALS!${hist_last}${H['revenue']}" if t == 0
                        else f"{p}{rows['rev']}")
            prev_nwc = (f"HISTORICALS!${hist_last}${H['nwc_balance']}" if t == 0
                        else f"{p}{rows['nwc']}")
            put(ws, f"{c}{rows['rev']}",
                f"={prev_rev}*(1+INPUTS!{c}${I['rev_growth_row']}{shift})", fmt=NUM)
            put(ws, f"{c}{rows['ebit']}",
                f"={c}{rows['rev']}*INPUTS!{c}${I['ebit_margin_row']}", fmt=NUM)
            put(ws, f"{c}{rows['nopat']}",
                f"={c}{rows['ebit']}*(1-INPUTS!$B$4)", fmt=NUM)
            put(ws, f"{c}{rows['da']}",
                f"={c}{rows['rev']}*INPUTS!{c}${I['da_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{rows['capex']}",
                f"={c}{rows['rev']}*INPUTS!{c}${I['capex_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{rows['nwc']}",
                f"={c}{rows['rev']}*INPUTS!{c}${I['nwc_pct_rev_row']}", fmt=NUM)
            put(ws, f"{c}{rows['dnwc']}", f"={c}{rows['nwc']}-{prev_nwc}", fmt=NUM)
            put(ws, f"{c}{rows['fcf']}",
                f"={c}{rows['nopat']}+{c}{rows['da']}-{c}{rows['capex']}"
                f"-{c}{rows['dnwc']}", fmt=NUM)
            put(ws, f"{c}{rows['pv']}",
                f"={c}{rows['fcf']}*VALUATION!{c}${V['discount_factor']}", fmt=NUM)

        ws.cell(row=rows["ev"], column=1, value="    Enterprise value").font = F_NOTE
        put(ws, f"B{rows['ev']}",
            f"=SUM(B{rows['pv']}:{flast}{rows['pv']})"
            f"+{flast}{rows['fcf']}*(1+INPUTS!$B$11)"
            f"/(INPUTS!$B$10-INPUTS!$B$11)*VALUATION!${flast}${V['discount_factor']}",
            fmt=NUM)
        ws.cell(row=rows["vps"], column=1, value="    Value per share").font = F_NOTE
        put(ws, f"B{rows['vps']}",
            f"=(B{rows['ev']}-INPUTS!$B$13)/INPUTS!$B$14", fmt=CUR)

    widths(ws, {"A": 34, "B": 15, "C": 16})
    for k in range(nf):
        ws.column_dimensions[col(2 + k)].width = 13
    return ws, S


# --------------------------------------------------------------------------
# football field
# --------------------------------------------------------------------------


def build_football(wb, spec, rel_map, scen_map, cells):
    ws = wb.create_sheet("FOOTBALL_FIELD")
    R, S, V = rel_map, scen_map, cells["VALUATION"]
    market = spec.get("market", {})

    ws["A1"] = "FOOTBALL FIELD — every method's range on one scale"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("A single number implies a precision none of these methods has. "
                "Where the ranges overlap is the useful part.")
    ws["A2"].font = F_NOTE
    head(ws, FF["header"], ["Method", "Low", "High", "Width", "Base (for the chart)"])

    rows = [
        ("Comps — EV/Revenue", f"=RELATIVE_PRICING!L{R['implied_p25']}",
         f"=RELATIVE_PRICING!L{R['implied_p75']}"),
        ("Comps — EV/EBITDA", f"=RELATIVE_PRICING!M{R['implied_p25']}",
         f"=RELATIVE_PRICING!M{R['implied_p75']}"),
        ("Comps — P/E", f"=RELATIVE_PRICING!N{R['implied_p25']}",
         f"=RELATIVE_PRICING!N{R['implied_p75']}"),
        ("DCF — bear to bull", f"=SCENARIOS!C{S['bear']}",
         f"=SCENARIOS!C{S['bull']}"),
        ("DCF — WACC/growth grid",
         "=MIN(SENSITIVITY!C5:G9)", "=MAX(SENSITIVITY!C5:G9)"),
        ("52-week trading range", None, None),
    ]
    for k, (lab, lo, hi) in enumerate(rows):
        r = FF["first"] + k
        label(ws, r, lab)
        if lo is None:
            put(ws, f"B{r}", market.get("week52_low", 0.0), font=F_INPUT, fmt=CUR)
            put(ws, f"C{r}", market.get("week52_high", 0.0), font=F_INPUT, fmt=CUR)
            note(ws, r, 6, market.get("source", ""))
        else:
            put(ws, f"B{r}", lo, fmt=CUR)
            put(ws, f"C{r}", hi, fmt=CUR)
        put(ws, f"D{r}", f"=C{r}-B{r}", fmt=CUR)
        put(ws, f"E{r}", f"=B{r}", fmt=CUR)

    r = FF["price"]
    label(ws, r, "Current share price")
    put(ws, f"B{r}", "=INPUTS!$B$15", fmt=CUR, fill=FILL_BAND)
    label(ws, r + 1, "DCF base case")
    put(ws, f"B{r + 1}", f"=VALUATION!{V['value_per_share']}", fmt=CUR,
        fill=FILL_OUT)
    note(ws, r + 1, 3,
         "Read this against the bands above, not on its own")

    last = FF["first"] + len(rows) - 1
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Implied value per share"
    chart.y_axis.title = "Value per share"
    base = Reference(ws, min_col=5, min_row=FF["first"], max_row=last)
    width = Reference(ws, min_col=4, min_row=FF["first"], max_row=last)
    cats = Reference(ws, min_col=1, min_row=FF["first"], max_row=last)
    chart.add_data(base, titles_from_data=False)
    chart.add_data(width, titles_from_data=False)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.noFill = True
    chart.height, chart.width = 9, 20
    ws.add_chart(chart, "G4")

    widths(ws, {"A": 32, "B": 14, "C": 14, "D": 14, "E": 20, "F": 44})
    return ws


# --------------------------------------------------------------------------
# beta regression
# --------------------------------------------------------------------------

BR = {"header": 3, "first": 4, "stats_col": 7}


def br_stats_rows():
    return {"observations": 4, "slope": 5, "intercept": 6, "r_squared": 7,
            "frequency": 8, "blume": 9}


def build_beta_regression(wb, spec):
    """The regression itself, not a number someone typed in.

    SLOPE, INTERCEPT, RSQ and COUNT are computed in the sheet over the return
    columns, so the beta can be re-derived by anyone holding the workbook.
    """
    reg = spec["wacc_build"]["regression_beta"]
    p = reg["prices"]
    n = len(p["stock"])
    dates = p.get("dates") or [""] * n
    S = br_stats_rows()
    first, last = BR["first"], BR["first"] + n - 1
    ws = wb.create_sheet("BETA_REGRESSION")

    ws["A1"] = "BETA REGRESSION — stock returns on index returns"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue columns are the price series. The beta is the slope of "
                "the fit, and R² says how much of the share's movement the "
                "index actually explains.")
    ws["A2"].font = F_NOTE
    head(ws, BR["header"], ["Date", "Stock price", "Stock return",
                            "Index price", "Index return"])

    for k in range(n):
        r = first + k
        put(ws, f"A{r}", dates[k], font=F_INPUT)
        put(ws, f"B{r}", p["stock"][k], font=F_INPUT, fmt=CUR)
        put(ws, f"D{r}", p["index"][k], font=F_INPUT, fmt=NUM)
        if k:
            put(ws, f"C{r}", f"=B{r}/B{r - 1}-1", fmt=PCT2)
            put(ws, f"E{r}", f"=D{r}/D{r - 1}-1", fmt=PCT2)

    y, x = f"C{first + 1}:C{last}", f"E{first + 1}:E{last}"
    c = col(BR["stats_col"])
    v = col(BR["stats_col"] + 1)
    ws.cell(row=S["observations"] - 1, column=BR["stats_col"],
            value="Regression output").font = F_TITLE
    for row, lab, formula, fmt in [
        (S["observations"], "Observations", f"=COUNT({y})", "0"),
        (S["slope"], "Slope (raw beta)", f"=SLOPE({y},{x})", "0.0000"),
        (S["intercept"], "Intercept (alpha)", f"=INTERCEPT({y},{x})", "0.00000"),
        (S["r_squared"], "R²", f"=RSQ({y},{x})", "0.0000"),
    ]:
        cell = ws.cell(row=row, column=BR["stats_col"], value=lab)
        cell.font = F_SUB
        put(ws, f"{v}{row}", formula, fmt=fmt,
            fill=FILL_OUT if row == S["slope"] else None)
    ws.cell(row=S["frequency"], column=BR["stats_col"],
            value="Frequency").font = F_SUB
    put(ws, f"{v}{S['frequency']}", reg.get("frequency", "weekly"), font=F_INPUT)
    note(ws, S["observations"], BR["stats_col"] + 2, reg.get("source", ""))
    note(ws, S["r_squared"], BR["stats_col"] + 2,
         "A low R² is a reason to lean on the peer-derived beta instead of this one")

    widths(ws, {"A": 13, "B": 13, "C": 13, "D": 14, "E": 13,
                col(BR["stats_col"]): 20, v: 14,
                col(BR["stats_col"] + 2): 56})
    ws.freeze_panes = "A4"
    return ws, S


# --------------------------------------------------------------------------
# segment revenue build
# --------------------------------------------------------------------------


def seg_rows(ns):
    r = {"header": 3, "first": 4}
    r["last"] = 3 + ns
    r["total"] = r["last"] + 1
    r["reported"] = r["total"] + 1
    r["check"] = r["total"] + 2
    r["implied"] = r["total"] + 3
    g = r["implied"] + 2
    r["growth_header"] = g
    r["growth_first"] = g + 1
    r["growth_last"] = g + ns
    m = r["growth_last"] + 2
    r["mix_header"] = m
    r["mix_first"] = m + 1
    r["mix_last"] = m + ns
    return r


def build_segments(wb, spec, hist_rows):
    """Revenue built from the segments up, so group growth is a consequence."""
    segs = spec["segments"]
    ns = len(segs)
    nh, nf = core.n_hist(spec), core.n_fcst(spec)
    S = seg_rows(ns)
    H = hist_rows
    hl = col(1 + nh)
    ws = wb.create_sheet("SEGMENTS")

    ws["A1"] = "SEGMENTS — the top line, built up"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Each segment on its own growth path. Group growth is then a "
                "consequence of the mix, not an assumption: a shift toward a "
                "faster segment lifts it without anyone asserting that it does.")
    ws["A2"].font = F_NOTE
    head(ws, S["header"], ["Segment"] + [str(y) for y in core.hist_years(spec)]
         + [str(y) for y in core.fcst_years(spec)])
    for k in range(nf):
        ws.cell(row=S["header"], column=2 + nh + k).fill = FILL_BAND

    for k, sg in enumerate(segs):
        r = S["first"] + k
        label(ws, r, sg["name"])
        for t in range(nh):
            put(ws, f"{col(2 + t)}{r}", sg["historical_revenue"][t],
                font=F_INPUT, fmt=NUM)
        for t in range(nf):
            c = col(2 + nh + t)
            prev = col(1 + nh + t) if t else hl
            put(ws, f"{c}{r}",
                f"={prev}{r}*(1+{c}{S['growth_first'] + k})", fmt=NUM)

    label(ws, S["total"], "Total revenue")
    label(ws, S["reported"], "Reported revenue")
    label(ws, S["check"], "Check (segments − reported)")
    label(ws, S["implied"], "Implied group growth")
    for t in range(nh + nf):
        c = col(2 + t)
        put(ws, f"{c}{S['total']}",
            f"=SUM({c}{S['first']}:{c}{S['last']})", fmt=NUM,
            fill=FILL_BAND if t >= nh else None)
        if t < nh:
            put(ws, f"{c}{S['reported']}", f"=HISTORICALS!{c}{H['revenue']}", fmt=NUM)
            put(ws, f"{c}{S['check']}",
                f"={c}{S['total']}-{c}{S['reported']}", fmt=NUM)
        else:
            prev = col(1 + t)
            put(ws, f"{c}{S['implied']}",
                f"={c}{S['total']}/{prev}{S['total']}-1", fmt=PCT,
                fill=FILL_OUT)
    note(ws, S["check"], 2 + nh + nf,
         "Must be zero in every reported year, or a segment is missing")

    ws.cell(row=S["growth_header"], column=1,
            value="Growth assumptions by segment").font = F_TITLE
    ws.cell(row=S["growth_header"], column=2 + nh + nf,
            value="Basis / source").font = F_SUB
    for k, sg in enumerate(segs):
        r = S["growth_first"] + k
        label(ws, r, sg["name"])
        for t in range(nf):
            put(ws, f"{col(2 + nh + t)}{r}", sg["growth"][t], font=F_INPUT, fmt=PCT)
        note(ws, r, 2 + nh + nf,
             f"{sg.get('basis', '')} — {sg.get('source', '')}")

    ws.cell(row=S["mix_header"], column=1,
            value="Share of revenue").font = F_TITLE
    for k, sg in enumerate(segs):
        r = S["mix_first"] + k
        label(ws, r, sg["name"])
        for t in range(nh + nf):
            c = col(2 + t)
            put(ws, f"{c}{r}",
                f"={c}{S['first'] + k}/{c}{S['total']}", fmt=PCT)

    year_widths(ws, nh + nf, label_w=30, year_w=11, extra=52)
    ws.freeze_panes = "B4"
    return ws, S
