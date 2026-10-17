"""Shared audit machinery: the result table and the LibreOffice recalculation.

Every model type is audited the same way -- recalculate the real workbook,
then compare it against an independently written Python engine. Nothing here
trusts the workbook's own cached values.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile

from openpyxl import load_workbook

TOL = 1e-6
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!")

RECALC_PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" \
xmlns:xs="http://www.w3.org/2001/XMLSchema" \
xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">\
<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">\
<prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


class Audit:
    def __init__(self, title="Financopedia model audit"):
        self.title = title
        self.rows = []

    def check(self, name, ok, detail=""):
        self.rows.append((name, bool(ok), detail))
        return ok

    def info(self, name, detail):
        self.rows.append((name, None, detail))

    @property
    def failures(self):
        return [r for r in self.rows if r[1] is False]

    def report(self):
        width = max(len(r[0]) for r in self.rows) + 2
        print(f"\n{self.title}")
        print("=" * (width + 44))
        for name, ok, detail in self.rows:
            mark = "INFO" if ok is None else ("PASS" if ok else "FAIL")
            print(f"  [{mark}] {name:<{width}} {detail}")
        total = len([r for r in self.rows if r[1] is not None])
        passed = total - len(self.failures)
        print("=" * (width + 44))
        print(f"  {passed}/{total} checks passed")
        if self.failures:
            print("\n  DO NOT DELIVER. Fix these and re-run:")
            for name, _, detail in self.failures:
                print(f"    - {name}: {detail}")
        return 0 if not self.failures else 1


def recalculate(path):
    """Round-trip through LibreOffice so formulas acquire computed values.

    LibreOffice does NOT recalculate xlsx on load by default -- it trusts the
    cached values, and openpyxl writes none, so an unconfigured conversion
    returns a workbook full of blanks. The profile forces a full
    recalculation. Input and output must live in different directories:
    converting xlsx to xlsx in place leaves the original untouched.
    """
    tmp = tempfile.mkdtemp(prefix="financopedia_audit_")
    profile = os.path.join(tmp, "loprofile")
    os.makedirs(os.path.join(profile, "user"), exist_ok=True)
    with open(os.path.join(profile, "user", "registrymodifications.xcu"),
              "w", encoding="utf-8") as fh:
        fh.write(RECALC_PROFILE)

    indir, outdir = os.path.join(tmp, "in"), os.path.join(tmp, "out")
    os.makedirs(indir, exist_ok=True)
    os.makedirs(outdir, exist_ok=True)
    src = os.path.join(indir, os.path.basename(path))
    shutil.copy(path, src)

    proc = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}",
         "--headless", "--norestore", "--convert-to", "xlsx",
         "--outdir", outdir, src],
        capture_output=True, text=True, timeout=600,
    )
    out = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    if not os.path.exists(out):
        raise RuntimeError("LibreOffice could not recalculate the workbook.\n"
                           f"{proc.stdout}\n{proc.stderr}")
    return out, tmp


def close(a, b, tol=TOL):
    if a is None or b is None:
        return False
    return abs(a - b) / max(abs(a), abs(b), 1.0) < tol


def formula_cells_only(a, wb_f, sheets, exclude_sheets=()):
    """No hardcoded constants on calculation sheets."""
    stray = []
    for sheet in sheets:
        if sheet in exclude_sheets or sheet not in wb_f.sheetnames:
            continue
        for row in wb_f[sheet].iter_rows(min_col=2):
            for cell in row:
                if isinstance(cell.value, (int, float)) and cell.value is not None:
                    stray.append(f"{sheet}!{cell.coordinate}={cell.value}")
    return a.check("No hardcoded numbers outside INPUTS/HISTORICALS", not stray,
                   f"{len(stray)} stray constants: {stray[:6]}" if stray
                   else f"every value cell on {', '.join(sheets)} is a formula")


def no_error_values(a, wbv):
    errs = []
    for sheet in wbv.sheetnames:
        for row in wbv[sheet].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in ERRORS:
                    errs.append(f"{sheet}!{cell.coordinate}={cell.value}")
    return a.check("No broken references or error values", not errs,
                   f"{len(errs)} errors: {errs[:5]}" if errs
                   else "no #REF!/#DIV0!/#VALUE!")


def workbook_checks_tab(a, wbv):
    ws = wbv["CHECKS"]
    results, failed = [], []
    for r in range(4, 60):
        v = ws.cell(row=r, column=2).value
        if v in ("PASS", "FAIL", "REVIEW"):
            results.append(v)
            if v == "FAIL":
                failed.append(str(ws.cell(row=r, column=1).value))
    return a.check("In-workbook CHECKS tab all pass", results and not failed,
                   f"FAILED: {failed}" if failed
                   else f"{results.count('PASS')}/{len(results)} PASS"
                   + (f", {results.count('REVIEW')} REVIEW"
                      if "REVIEW" in results else ""))


def compare(a, name, pairs, fmt="{:,.2f}"):
    """pairs: list of (label, excel_value, python_value)."""
    bad = []
    for lab, got, want in pairs:
        if not close(got, want):
            bad.append(f"{lab}: excel=" + ("blank" if got is None
                                           else fmt.format(got))
                       + f" python={fmt.format(want)}")
    return a.check(name, not bad, "; ".join(bad[:6]) if bad
                   else " · ".join(f"{lab} {fmt.format(want)}"
                                   for lab, _, want in pairs[:6]))


def getv(wbv, sheet, addr):
    v = wbv[sheet][addr].value
    return float(v) if isinstance(v, (int, float)) else None


def open_pair(path):
    """Return (formula workbook, recalculated value workbook, tmpdir)."""
    wb_f = load_workbook(path, data_only=False)
    recalc, tmp = recalculate(path)
    wb_v = load_workbook(recalc, data_only=True)
    return wb_f, wb_v, tmp
