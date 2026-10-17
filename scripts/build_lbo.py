#!/usr/bin/env python3
"""LBO workbook: sources and uses, tranched debt with a cash sweep, exit returns."""

from __future__ import annotations

from openpyxl import Workbook

import dcf_core as core
import lbo_core as lbo
from xl_style import (CUR, F_ESTIMATE, F_INPUT, F_NOTE, F_SUB, F_TITLE,
                      FILL_BAND, FILL_EST, FILL_OUT, MULT, NUM, PCT,
                      checks_sheet, col, head, label, note, put, readme_sheet,
                      widths, year_widths)

GRID = 5

LBOC = {
    "INPUTS": {
        "entry_multiple": "B4", "exit_multiple": "B5",
        "transaction_fees_pct": "B6", "financing_fees_pct": "B7",
        "rollover_equity": "B8", "tax_rate": "B9", "cash_sweep_pct": "B10",
        "min_cash": "B11", "opening_cash": "B12", "interest_rate_cash": "B13",
        "entry_step": "B14", "exit_step": "B15",
        "rev_growth": 18, "ebitda_margin": 19, "da_pct_rev": 20,
        "capex_pct_rev": 21, "nwc_pct_rev": 22,
        "tranche_start": 25,
    },
    "HISTORICALS": {
        "revenue": 4, "ebitda": 5, "capex": 6, "nwc_balance": 7,
        "margin": 8, "cagr": 11, "margin_avg": 12, "margin_min": 13,
    },
    "PROJECTION": {
        "revenue": 4, "ebitda": 5, "da": 6, "ebit": 7, "interest": 8,
        "interest_income": 9, "pbt": 10, "tax": 11, "net_income": 12,
        "capex": 13, "nwc": 14, "delta_nwc": 15, "fcf": 16,
    },
    "SENSITIVITY": {"x_row": 4, "y_start": 5, "grid_first_col": 3},
}


def su_rows(nt):
    """Sources and uses row map for nt tranches."""
    r = {"ltm_ebitda": 4, "entry_multiple": 5, "entry_ev": 6,
         "sources_header": 8, "tranche_first": 9}
    r["total_debt"] = 9 + nt
    r["rollover"] = r["total_debt"] + 1
    r["sponsor"] = r["total_debt"] + 2
    r["total_sources"] = r["total_debt"] + 3
    r["uses_header"] = r["total_debt"] + 5
    r["purchase"] = r["total_debt"] + 6
    r["transaction_fees"] = r["total_debt"] + 7
    r["financing_fees"] = r["total_debt"] + 8
    r["cash_to_bs"] = r["total_debt"] + 9
    r["total_uses"] = r["total_debt"] + 10
    r["check"] = r["total_debt"] + 12
    r["entry_net_debt"] = r["total_debt"] + 13
    r["entry_equity"] = r["total_debt"] + 14
    r["entry_leverage"] = r["total_debt"] + 15
    return r


def debt_rows(nt):
    """Debt schedule row map: a block per tranche, then the totals block."""
    rows = {"blocks": []}
    for k in range(nt):
        base = 4 + k * 6
        rows["blocks"].append({
            "header": base, "opening": base + 1, "mandatory": base + 2,
            "sweep": base + 3, "closing": base + 4, "interest": base + 5})
    t = 4 + nt * 6 + 1
    for n, key in enumerate([
            "total_opening", "total_mandatory", "total_sweep", "total_closing",
            "total_interest", "opening_cash", "fcf", "cash_before_sweep",
            "cash_available", "closing_cash", "net_debt", "leverage"]):
        rows[key] = t + n
    return rows


RET = {"exit_ebitda": 4, "exit_multiple": 5, "exit_ev": 6, "exit_net_debt": 7,
       "exit_equity": 8, "sponsor": 9, "money_multiple": 10, "years": 11,
       "irr": 12, "attr_header": 14, "attr_ebitda": 15, "attr_multiple": 16,
       "attr_debt": 17, "attr_total": 18, "attr_check": 19}


def build_inputs(wb, spec):
    ws = wb.create_sheet("INPUTS")
    i = spec["inputs"]
    basis, est = i.get("basis", {}), set(i.get("estimates", []))
    nf = core.n_fcst(spec)
    bcol = max(2 + nf, 5)
    I = LBOC["INPUTS"]

    ws["A1"] = "INPUTS — deal terms, capital structure and operating case"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Debt is sized in turns of LTM EBITDA, so the structure does not "
                "move when the entry multiple does.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Assumption", "Value", "Unit"])
    ws.cell(row=3, column=bcol, value="Basis / source").font = F_SUB

    scalars = [
        ("entry_multiple", "Entry EV/EBITDA multiple", MULT, "x"),
        ("exit_multiple", "Exit EV/EBITDA multiple", MULT, "x"),
        ("transaction_fees_pct", "Transaction fees (% of EV)", PCT, "%"),
        ("financing_fees_pct", "Financing fees (% of new debt)", PCT, "%"),
        ("rollover_equity", "Rollover equity", NUM, spec["meta"]["units"]),
        ("tax_rate", "Tax rate", PCT, "%"),
        ("cash_sweep_pct", "Cash sweep (% of surplus)", PCT, "%"),
        ("min_cash", "Minimum cash balance", NUM, spec["meta"]["units"]),
        ("opening_cash", "Cash left on the balance sheet at close", NUM,
         spec["meta"]["units"]),
        ("interest_rate_cash", "Interest earned on cash", PCT, "%"),
        ("entry_step", "Sensitivity: entry multiple step", MULT, "x"),
        ("exit_step", "Sensitivity: exit multiple step", MULT, "x"),
    ]
    for key, lab, fmt, unit in scalars:
        r = int(I[key][1:])
        label(ws, r, lab)
        is_est = key in est
        put(ws, I[key], i[key], font=F_ESTIMATE if is_est else F_INPUT, fmt=fmt,
            fill=FILL_EST if is_est else None)
        ws.cell(row=r, column=3, value=unit)
        note(ws, r, bcol, basis.get(key, ""))

    head(ws, 17, ["Operating driver"] + [f"Y{k+1}" for k in range(nf)])
    ws.cell(row=17, column=bcol, value="Basis / source").font = F_SUB
    for key, lab, fmt in [("rev_growth", "Revenue growth", PCT),
                          ("ebitda_margin", "EBITDA margin", PCT),
                          ("da_pct_rev", "D&A (% of revenue)", PCT),
                          ("capex_pct_rev", "Capex (% of revenue)", PCT),
                          ("nwc_pct_rev", "Working capital (% of revenue)", PCT)]:
        r = I[key]
        label(ws, r, lab)
        is_est = key in est
        for t in range(nf):
            put(ws, f"{col(2 + t)}{r}", i[key][t],
                font=F_ESTIMATE if is_est else F_INPUT, fmt=fmt,
                fill=FILL_EST if is_est else None)
        note(ws, r, bcol, basis.get(key, ""))

    head(ws, 24, ["Debt tranche (in priority order)", "Turns of EBITDA",
                  "Rate", "Annual amortisation", "In the sweep?"])
    ws.cell(row=24, column=6, value="Basis / source").font = F_SUB
    for k, tr in enumerate(spec["debt_tranches"]):
        r = I["tranche_start"] + k
        label(ws, r, tr["name"])
        put(ws, f"B{r}", tr["turns"], font=F_INPUT, fmt=MULT)
        put(ws, f"C{r}", tr["rate"], font=F_INPUT, fmt=PCT)
        put(ws, f"D{r}", tr["amort_pct"], font=F_INPUT, fmt=PCT)
        put(ws, f"E{r}", 1 if tr.get("sweep", True) else 0, font=F_INPUT, fmt="0")
        note(ws, r, 6, tr.get("basis", ""))

    year_widths(ws, nf, label_w=38, year_w=12)
    widths(ws, {"C": 13, "D": 20, "E": 14, "F": 52, col(bcol): 56})
    return ws


def build_historicals(wb, spec):
    ws = wb.create_sheet("HISTORICALS")
    h, src = spec["historicals"], spec["historicals"].get("sources", {})
    nh = core.n_hist(spec)
    scol, last = 2 + nh, col(1 + nh)
    H = LBOC["HISTORICALS"]

    ws["A1"] = f"HISTORICALS — {nh} reported years"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("The last column is LTM: it sizes the debt and prices the "
                "entry, so it carries more weight than any other figure here.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Metric"] + [str(y) for y in h["years"]])
    ws.cell(row=3, column=scol, value="Source (filing, date)").font = F_SUB

    for key, lab in [("revenue", "Revenue"), ("ebitda", "EBITDA"),
                     ("capex", "Capex"),
                     ("nwc_balance", "Net working capital (balance)")]:
        r = H[key]
        label(ws, r, lab)
        for t in range(nh):
            put(ws, f"{col(2 + t)}{r}", h[key][t], font=F_INPUT, fmt=NUM)
        note(ws, r, scol, src.get(key, ""))

    label(ws, H["margin"], "EBITDA margin")
    for t in range(nh):
        c = col(2 + t)
        put(ws, f"{c}{H['margin']}", f"={c}{H['ebitda']}/{c}{H['revenue']}", fmt=PCT)

    ws.cell(row=10, column=1, value=f"What {nh} years of history says").font = F_TITLE
    for r, lab, formula, fmt, why in [
        (H["cagr"], "Revenue CAGR over the period",
         f"=(({last}{H['revenue']}/B{H['revenue']})^(1/{nh - 1}))-1", PCT,
         "A leveraged structure is paid down out of growth; compare with the case"),
        (H["margin_avg"], "EBITDA margin — average",
         f"=AVERAGE(B{H['margin']}:{last}{H['margin']})", PCT, ""),
        (H["margin_min"], "EBITDA margin — worst year",
         f"=MIN(B{H['margin']}:{last}{H['margin']})", PCT,
         "Test the structure against this, not the average — covenants break in bad years"),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt, fill=FILL_BAND)
        if why:
            note(ws, r, 3, why)

    year_widths(ws, nh, label_w=34, year_w=12, extra=56)
    return ws


def build_sources_uses(wb, spec):
    ws = wb.create_sheet("SOURCES_USES")
    nh, nt = core.n_hist(spec), len(spec["debt_tranches"])
    HL = col(1 + nh)
    H, I = LBOC["HISTORICALS"], LBOC["INPUTS"]
    S = su_rows(nt)

    ws["A1"] = "SOURCES AND USES — the entry"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Sponsor equity is the plug: it is whatever the uses are not funded by."
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Item", "Amount", "Turns of EBITDA"])

    label(ws, S["ltm_ebitda"], "LTM EBITDA")
    put(ws, f"B{S['ltm_ebitda']}", f"=HISTORICALS!{HL}{H['ebitda']}", fmt=NUM)
    label(ws, S["entry_multiple"], "Entry multiple")
    put(ws, f"B{S['entry_multiple']}", "=INPUTS!$B$4", fmt=MULT)
    label(ws, S["entry_ev"], "Entry enterprise value")
    put(ws, f"B{S['entry_ev']}",
        f"=B{S['ltm_ebitda']}*B{S['entry_multiple']}", fmt=NUM, fill=FILL_BAND)

    ws.cell(row=S["sources_header"], column=1, value="SOURCES").font = F_TITLE
    for k, tr in enumerate(spec["debt_tranches"]):
        r = S["tranche_first"] + k
        label(ws, r, tr["name"])
        put(ws, f"B{r}",
            f"=INPUTS!B{I['tranche_start'] + k}*$B${S['ltm_ebitda']}", fmt=NUM)
        put(ws, f"C{r}", f"=B{r}/$B${S['ltm_ebitda']}", fmt=MULT)
    label(ws, S["total_debt"], "Total new debt")
    put(ws, f"B{S['total_debt']}",
        f"=SUM(B{S['tranche_first']}:B{S['total_debt'] - 1})", fmt=NUM)
    put(ws, f"C{S['total_debt']}",
        f"=B{S['total_debt']}/$B${S['ltm_ebitda']}", fmt=MULT)
    label(ws, S["rollover"], "Rollover equity")
    put(ws, f"B{S['rollover']}", "=INPUTS!$B$8", fmt=NUM)
    label(ws, S["sponsor"], "Sponsor equity (plug)")
    put(ws, f"B{S['sponsor']}",
        f"=B{S['total_uses']}-B{S['total_debt']}-B{S['rollover']}",
        fmt=NUM, fill=FILL_OUT)
    label(ws, S["total_sources"], "Total sources")
    put(ws, f"B{S['total_sources']}",
        f"=B{S['total_debt']}+B{S['rollover']}+B{S['sponsor']}", fmt=NUM)

    ws.cell(row=S["uses_header"], column=1, value="USES").font = F_TITLE
    label(ws, S["purchase"], "Purchase of enterprise")
    put(ws, f"B{S['purchase']}", f"=B{S['entry_ev']}", fmt=NUM)
    label(ws, S["transaction_fees"], "Transaction fees")
    put(ws, f"B{S['transaction_fees']}",
        f"=B{S['entry_ev']}*INPUTS!$B$6", fmt=NUM)
    label(ws, S["financing_fees"], "Financing fees")
    put(ws, f"B{S['financing_fees']}",
        f"=B{S['total_debt']}*INPUTS!$B$7", fmt=NUM)
    label(ws, S["cash_to_bs"], "Cash to the balance sheet")
    put(ws, f"B{S['cash_to_bs']}", "=INPUTS!$B$12", fmt=NUM)
    label(ws, S["total_uses"], "Total uses")
    put(ws, f"B{S['total_uses']}",
        f"=SUM(B{S['purchase']}:B{S['cash_to_bs']})", fmt=NUM)

    label(ws, S["check"], "Check: sources − uses")
    put(ws, f"B{S['check']}",
        f"=B{S['total_sources']}-B{S['total_uses']}", fmt=NUM, fill=FILL_OUT)
    label(ws, S["entry_net_debt"], "Net debt at close")
    put(ws, f"B{S['entry_net_debt']}",
        f"=B{S['total_debt']}-INPUTS!$B$12", fmt=NUM)
    label(ws, S["entry_equity"], "Entry equity value (EV − net debt)")
    put(ws, f"B{S['entry_equity']}",
        f"=B{S['entry_ev']}-B{S['entry_net_debt']}", fmt=NUM)
    label(ws, S["entry_leverage"], "Opening leverage (net debt / EBITDA)")
    put(ws, f"B{S['entry_leverage']}",
        f"=B{S['entry_net_debt']}/B{S['ltm_ebitda']}", fmt=MULT, fill=FILL_BAND)

    widths(ws, {"A": 42, "B": 16, "C": 18})
    return ws


def build_projection(wb, spec):
    ws = wb.create_sheet("PROJECTION")
    nh, nf, nt = core.n_hist(spec), core.n_fcst(spec), len(spec["debt_tranches"])
    HL = col(1 + nh)
    H, I, P = LBOC["HISTORICALS"], LBOC["INPUTS"], LBOC["PROJECTION"]
    DR = debt_rows(nt)

    ws["A1"] = "PROJECTION — the operating case under the new structure"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Interest comes from the DEBT tab and is charged on opening "
                "balances, so nothing here is circular.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Metric"] + [str(y) for y in core.fcst_years(spec)])

    for key, lab in [("revenue", "Revenue"), ("ebitda", "EBITDA"),
                     ("da", "D&A"), ("ebit", "EBIT"),
                     ("interest", "Interest expense"),
                     ("interest_income", "Interest income"),
                     ("pbt", "Profit before tax"), ("tax", "Tax"),
                     ("net_income", "Net income"), ("capex", "Capex"),
                     ("nwc", "Working capital (balance)"),
                     ("delta_nwc", "Change in working capital"),
                     ("fcf", "Free cash flow before debt service")]:
        label(ws, P[key], lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        prev_rev = (f"HISTORICALS!${HL}${H['revenue']}" if t == 0
                    else f"{p}{P['revenue']}")
        prev_nwc = (f"HISTORICALS!${HL}${H['nwc_balance']}" if t == 0
                    else f"{p}{P['nwc']}")
        put(ws, f"{c}{P['revenue']}",
            f"={prev_rev}*(1+INPUTS!{c}${I['rev_growth']})", fmt=NUM)
        put(ws, f"{c}{P['ebitda']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['ebitda_margin']}", fmt=NUM,
            fill=FILL_OUT)
        put(ws, f"{c}{P['da']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['da_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{P['ebit']}", f"={c}{P['ebitda']}-{c}{P['da']}", fmt=NUM)
        put(ws, f"{c}{P['interest']}", f"=DEBT!{c}{DR['total_interest']}", fmt=NUM)
        put(ws, f"{c}{P['interest_income']}",
            f"=DEBT!{c}{DR['opening_cash']}*INPUTS!$B$13", fmt=NUM)
        put(ws, f"{c}{P['pbt']}",
            f"={c}{P['ebit']}-{c}{P['interest']}+{c}{P['interest_income']}", fmt=NUM)
        put(ws, f"{c}{P['tax']}", f"=MAX({c}{P['pbt']},0)*INPUTS!$B$9", fmt=NUM)
        put(ws, f"{c}{P['net_income']}", f"={c}{P['pbt']}-{c}{P['tax']}", fmt=NUM)
        put(ws, f"{c}{P['capex']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['capex_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{P['nwc']}",
            f"={c}{P['revenue']}*INPUTS!{c}${I['nwc_pct_rev']}", fmt=NUM)
        put(ws, f"{c}{P['delta_nwc']}", f"={c}{P['nwc']}-{prev_nwc}", fmt=NUM)
        put(ws, f"{c}{P['fcf']}",
            f"={c}{P['net_income']}+{c}{P['da']}-{c}{P['capex']}"
            f"-{c}{P['delta_nwc']}", fmt=NUM, fill=FILL_OUT)

    year_widths(ws, nf, label_w=38)
    return ws


def build_debt(wb, spec):
    ws = wb.create_sheet("DEBT")
    nf, nt = core.n_fcst(spec), len(spec["debt_tranches"])
    I, P = LBOC["INPUTS"], LBOC["PROJECTION"]
    S, DR = su_rows(nt), debt_rows(nt)

    ws["A1"] = "DEBT SCHEDULE — mandatory amortisation then cash sweep"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Tranches are swept in the order listed. Each year's interest is "
                "charged on that year's opening balance.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Line"] + [str(y) for y in core.fcst_years(spec)])

    for k, tr in enumerate(spec["debt_tranches"]):
        b = DR["blocks"][k]
        ws.cell(row=b["header"], column=1, value=tr["name"]).font = F_TITLE
        for key, lab in [("opening", "Opening balance"),
                         ("mandatory", "Mandatory amortisation"),
                         ("sweep", "Cash sweep"),
                         ("closing", "Closing balance"),
                         ("interest", "Interest expense")]:
            label(ws, b[key], "    " + lab, font=F_SUB)

    for key, lab in [("total_opening", "Total opening debt"),
                     ("total_mandatory", "Total mandatory amortisation"),
                     ("total_sweep", "Total cash sweep"),
                     ("total_closing", "Total closing debt"),
                     ("total_interest", "Total interest expense"),
                     ("opening_cash", "Opening cash"),
                     ("fcf", "Free cash flow before debt service"),
                     ("cash_before_sweep", "Cash before sweep"),
                     ("cash_available", "Cash available for sweep"),
                     ("closing_cash", "Closing cash"),
                     ("net_debt", "Net debt"),
                     ("leverage", "Net debt / EBITDA")]:
        label(ws, DR[key], lab)

    for t in range(nf):
        c, p = col(2 + t), col(1 + t)
        # opening balances and mandatory amortisation first
        for k in range(nt):
            b = DR["blocks"][k]
            orig = f"SOURCES_USES!$B${S['tranche_first'] + k}"
            opening = orig if t == 0 else f"{p}{b['closing']}"
            put(ws, f"{c}{b['opening']}", f"={opening}", fmt=NUM)
            put(ws, f"{c}{b['mandatory']}",
                f"=MIN(INPUTS!$D${I['tranche_start'] + k}*{orig},{c}{b['opening']})",
                fmt=NUM)
            put(ws, f"{c}{b['interest']}",
                f"={c}{b['opening']}*INPUTS!$C${I['tranche_start'] + k}", fmt=NUM)

        first, last_b = DR["blocks"][0], DR["blocks"][nt - 1]
        put(ws, f"{c}{DR['total_opening']}",
            "=" + "+".join(f"{c}{DR['blocks'][k]['opening']}" for k in range(nt)),
            fmt=NUM)
        put(ws, f"{c}{DR['total_mandatory']}",
            "=" + "+".join(f"{c}{DR['blocks'][k]['mandatory']}" for k in range(nt)),
            fmt=NUM)
        put(ws, f"{c}{DR['total_interest']}",
            "=" + "+".join(f"{c}{DR['blocks'][k]['interest']}" for k in range(nt)),
            fmt=NUM)
        open_cash = "INPUTS!$B$12" if t == 0 else f"{p}{DR['closing_cash']}"
        put(ws, f"{c}{DR['opening_cash']}", f"={open_cash}", fmt=NUM)
        put(ws, f"{c}{DR['fcf']}", f"=PROJECTION!{c}{P['fcf']}", fmt=NUM)
        put(ws, f"{c}{DR['cash_before_sweep']}",
            f"={c}{DR['opening_cash']}+{c}{DR['fcf']}-{c}{DR['total_mandatory']}",
            fmt=NUM)
        put(ws, f"{c}{DR['cash_available']}",
            f"=MAX(0,{c}{DR['cash_before_sweep']}-INPUTS!$B$11)*INPUTS!$B$10",
            fmt=NUM)

        # sweeps in priority order: each tranche takes what is left
        for k in range(nt):
            b = DR["blocks"][k]
            taken = ("" if k == 0 else
                     "-" + "-".join(f"{c}{DR['blocks'][j]['sweep']}"
                                    for j in range(k)))
            put(ws, f"{c}{b['sweep']}",
                f"=INPUTS!$E${I['tranche_start'] + k}*MIN({c}{b['opening']}"
                f"-{c}{b['mandatory']},MAX(0,{c}{DR['cash_available']}{taken}))",
                fmt=NUM)
            put(ws, f"{c}{b['closing']}",
                f"={c}{b['opening']}-{c}{b['mandatory']}-{c}{b['sweep']}", fmt=NUM)

        put(ws, f"{c}{DR['total_sweep']}",
            "=" + "+".join(f"{c}{DR['blocks'][k]['sweep']}" for k in range(nt)),
            fmt=NUM)
        put(ws, f"{c}{DR['total_closing']}",
            "=" + "+".join(f"{c}{DR['blocks'][k]['closing']}" for k in range(nt)),
            fmt=NUM, fill=FILL_OUT)
        put(ws, f"{c}{DR['closing_cash']}",
            f"={c}{DR['cash_before_sweep']}-{c}{DR['total_sweep']}", fmt=NUM)
        put(ws, f"{c}{DR['net_debt']}",
            f"={c}{DR['total_closing']}-{c}{DR['closing_cash']}", fmt=NUM)
        put(ws, f"{c}{DR['leverage']}",
            f"={c}{DR['net_debt']}/PROJECTION!{c}{P['ebitda']}", fmt=MULT,
            fill=FILL_BAND)

    year_widths(ws, nf, label_w=40)
    return ws


def build_returns(wb, spec):
    ws = wb.create_sheet("RETURNS")
    nf, nt = core.n_fcst(spec), len(spec["debt_tranches"])
    FL = col(1 + nf)
    P, S, DR = LBOC["PROJECTION"], su_rows(nt), debt_rows(nt)

    ws["A1"] = "RETURNS — exit and attribution"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Attribution is an identity, not an estimate: the three "
                "components sum exactly to the change in equity value.")
    ws["A2"].font = F_NOTE
    head(ws, 3, ["Item", "Value"])

    for r, lab, formula, fmt, fill in [
        (RET["exit_ebitda"], "Exit-year EBITDA",
         f"=PROJECTION!{FL}{P['ebitda']}", NUM, None),
        (RET["exit_multiple"], "Exit multiple", "=INPUTS!$B$5", MULT, None),
        (RET["exit_ev"], "Exit enterprise value",
         f"=B{RET['exit_ebitda']}*B{RET['exit_multiple']}", NUM, None),
        (RET["exit_net_debt"], "Less: net debt at exit",
         f"=DEBT!{FL}{DR['net_debt']}", NUM, None),
        (RET["exit_equity"], "Exit equity value",
         f"=B{RET['exit_ev']}-B{RET['exit_net_debt']}", NUM, FILL_OUT),
        (RET["sponsor"], "Sponsor equity at entry",
         f"=SOURCES_USES!B{S['sponsor']}", NUM, None),
        (RET["money_multiple"], "Money multiple",
         f"=B{RET['exit_equity']}/B{RET['sponsor']}", MULT, FILL_OUT),
        (RET["years"], "Holding period (years)",
         f"=COLUMNS(PROJECTION!B{P['ebitda']}:{FL}{P['ebitda']})", "0", None),
        (RET["irr"], "IRR",
         f"=IF(OR(B{RET['sponsor']}<=0,B{RET['exit_equity']}<=0),NA(),"
         f"B{RET['money_multiple']}^(1/B{RET['years']})-1)", PCT, FILL_OUT),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=fmt, fill=fill)

    ws.cell(row=RET["attr_header"], column=1,
            value="Where the value came from").font = F_TITLE
    for r, lab, formula, why in [
        (RET["attr_ebitda"], "EBITDA growth",
         f"=(B{RET['exit_ebitda']}-SOURCES_USES!B{S['ltm_ebitda']})*INPUTS!$B$4",
         "Growth in earnings, valued at the entry multiple"),
        (RET["attr_multiple"], "Multiple change",
         f"=(INPUTS!$B$5-INPUTS!$B$4)*B{RET['exit_ebitda']}",
         "Re-rating, applied to exit earnings. The component you do not control"),
        (RET["attr_debt"], "Debt paydown and cash generation",
         f"=SOURCES_USES!B{S['entry_net_debt']}-B{RET['exit_net_debt']}",
         "Net debt reduction over the hold"),
        (RET["attr_total"], "Total value created",
         f"=B{RET['attr_ebitda']}+B{RET['attr_multiple']}+B{RET['attr_debt']}", ""),
        (RET["attr_check"], "Check: total − (exit equity − entry equity)",
         f"=B{RET['attr_total']}-(B{RET['exit_equity']}"
         f"-SOURCES_USES!B{S['entry_equity']})",
         "Must be zero — the attribution is an identity"),
    ]:
        label(ws, r, lab)
        put(ws, f"B{r}", formula, fmt=NUM,
            fill=FILL_OUT if r == RET["attr_total"] else None)
        if why:
            note(ws, r, 3, why)

    widths(ws, {"A": 46, "B": 16, "C": 58})
    return ws


def build_sensitivity(wb, spec):
    """Entry multiple vs exit multiple, on IRR.

    Debt is sized in turns of EBITDA, so neither the schedule nor the cash
    flows move with either multiple -- only the sponsor cheque and the exit
    equity do. That makes every cell an exact closed form.
    """
    ws = wb.create_sheet("SENSITIVITY")
    nf, nt = core.n_fcst(spec), len(spec["debt_tranches"])
    S, DRr = su_rows(nt), LBOC["SENSITIVITY"]

    ws["A1"] = "SENSITIVITY — IRR: entry multiple vs exit multiple"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Both axes are driven off INPUTS by the step cells (B14/B15). "
                "Buying cheaply and selling dearly is the whole grid.")
    ws["A2"].font = F_NOTE
    ws.cell(row=DRr["x_row"] - 1, column=2,
            value="Entry ↓ / Exit →").font = F_SUB

    for k in range(GRID):
        c = col(DRr["grid_first_col"] + k)
        put(ws, f"{c}{DRr['x_row']}",
            f"=INPUTS!$B$5+(COLUMN({c}{DRr['x_row']})-"
            f"{DRr['grid_first_col'] + GRID // 2})*INPUTS!$B$15",
            font=F_SUB, fmt=MULT, fill=FILL_BAND)
    for j in range(GRID):
        r = DRr["y_start"] + j
        put(ws, f"B{r}",
            f"=INPUTS!$B$4+(ROW(B{r})-{DRr['y_start'] + GRID // 2})*INPUTS!$B$14",
            font=F_SUB, fmt=MULT, fill=FILL_BAND)

    for k in range(GRID):
        xc = col(DRr["grid_first_col"] + k)
        for j in range(GRID):
            r = DRr["y_start"] + j
            ev = f"$B{r}*SOURCES_USES!$B${S['ltm_ebitda']}"
            sponsor = (f"({ev}*(1+INPUTS!$B$6)+SOURCES_USES!$B${S['financing_fees']}"
                       f"+INPUTS!$B$12-SOURCES_USES!$B${S['total_debt']}"
                       f"-INPUTS!$B$8)")
            exit_eq = (f"({xc}${DRr['x_row']}*RETURNS!$B${RET['exit_ebitda']}"
                       f"-RETURNS!$B${RET['exit_net_debt']})")
            put(ws, f"{xc}{r}",
                f"=IF(OR({sponsor}<=0,{exit_eq}<=0),NA(),"
                f"({exit_eq}/{sponsor})^(1/RETURNS!$B${RET['years']})-1)",
                fmt=PCT,
                fill=FILL_OUT if (j == GRID // 2 and k == GRID // 2) else None)

    ws.cell(row=DRr["y_start"] + GRID + 1, column=2,
            value="Shaded centre cell = base case; it must equal the IRR on "
                  "RETURNS (CHECKS verifies this).").font = F_NOTE
    year_widths(ws, GRID + 1, label_w=24, year_w=13)
    return ws


def build_checks(wb, spec):
    nf, nt = core.n_fcst(spec), len(spec["debt_tranches"])
    FL = col(1 + nf)
    S, DR, DRr = su_rows(nt), debt_rows(nt), LBOC["SENSITIVITY"]
    P = LBOC["PROJECTION"]
    cc = col(DRr["grid_first_col"] + GRID // 2)
    cr = DRr["y_start"] + GRID // 2
    b0 = DR["blocks"][0]

    def rng(sheet, row):
        return f"{sheet}!B{row}:{FL}{row}"

    return checks_sheet(wb, [
        ("Sources equal uses",
         f'=IF(ABS(SOURCES_USES!B{S["check"]})<0.01,"PASS","FAIL")',
         "The entry funds itself — the sponsor cheque is the plug"),
        ("Attribution ties to value created",
         f'=IF(ABS(RETURNS!B{RET["attr_check"]})<0.01,"PASS","FAIL")',
         "EBITDA growth + multiple change + debt paydown = change in equity "
         "value, exactly"),
        ("No debt balance goes negative",
         f'=IF(MIN({rng("DEBT", DR["total_closing"])})>=-0.01,"PASS","FAIL")',
         "Repayment and sweep never exceed what is outstanding"),
        ("Cash never falls below the minimum",
         f'=IF(MIN({rng("DEBT", DR["closing_cash"])})>=INPUTS!$B$11-0.01,'
         f'"PASS","FAIL")',
         "The sweep leaves the minimum balance behind"),
        ("Debt roll-forward holds in year 1",
         f'=IF(ABS(DEBT!B{b0["closing"]}-(DEBT!B{b0["opening"]}'
         f'-DEBT!B{b0["mandatory"]}-DEBT!B{b0["sweep"]}))<0.01,"PASS","FAIL")',
         "Closing = opening - amortisation - sweep on the senior tranche"),
        ("Sweep never exceeds cash available",
         f'=IF(SUMPRODUCT(MAX({rng("DEBT", DR["total_sweep"])}'
         f'-{rng("DEBT", DR["cash_available"])}))<0.01,"PASS","FAIL")',
         "The sweep cannot spend money the business did not generate"),
        ("Interest charged on opening balances",
         f'=IF(ABS(DEBT!B{b0["interest"]}-DEBT!B{b0["opening"]}'
         f'*INPUTS!$C${LBOC["INPUTS"]["tranche_start"]})<0.01,"PASS","FAIL")',
         "No circular reference in the interest calculation"),
        ("Money multiple ties to the equity values",
         f'=IF(ABS(RETURNS!B{RET["money_multiple"]}-RETURNS!B{RET["exit_equity"]}'
         f'/RETURNS!B{RET["sponsor"]})<0.0001,"PASS","FAIL")',
         "MoM = exit equity / sponsor equity"),
        ("IRR is consistent with the money multiple",
         f'=IF(ABS((1+RETURNS!B{RET["irr"]})^RETURNS!B{RET["years"]}'
         f'-RETURNS!B{RET["money_multiple"]})<0.0001,"PASS","FAIL")',
         "(1+IRR)^years = MoM, with no interim distributions"),
        ("Sensitivity centre equals the headline IRR",
         f'=IF(ABS(SENSITIVITY!{cc}{cr}-RETURNS!B{RET["irr"]})<0.0001,'
         f'"PASS","FAIL")',
         "The grid is wired to the same deal"),
        ("Opening leverage at or below 7.0x",
         f'=IF(SOURCES_USES!B{S["entry_leverage"]}<=7,"PASS","REVIEW")',
         "Above this, the structure is unusual enough to need saying out loud"),
        ("EBITDA margin at or above the historical trough",
         f'=IF(MIN({rng("INPUTS", LBOC["INPUTS"]["ebitda_margin"])})'
         f'>=HISTORICALS!$B${LBOC["HISTORICALS"]["margin_min"]},"PASS","REVIEW")',
         "The case does not assume the company never has a bad year again"),
    ])


def build(spec, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    readme_sheet(wb, spec, "LBO model", [
        ("History", f"{core.n_hist(spec)} reported years; the last is LTM and "
                    "sizes both the debt and the entry price"),
        ("Hold", f"{core.n_fcst(spec)} years to exit"),
        ("Debt sizing", "In turns of LTM EBITDA, so the structure does not move "
                        "when the entry multiple does"),
        ("Interest", "Charged on opening balances — no iterative calculation"),
        ("Sweep", "Surplus cash above the minimum balance repays tranches in the "
                  "order they are listed"),
    ])
    build_inputs(wb, spec)
    build_historicals(wb, spec)
    build_sources_uses(wb, spec)
    build_projection(wb, spec)
    build_debt(wb, spec)
    build_returns(wb, spec)
    build_sensitivity(wb, spec)
    build_checks(wb, spec)
    wb.save(out_path)
    return out_path
